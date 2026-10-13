import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.38.xlsx").replace("0.18.xlsx","0.37.xlsx").replace("Version 0.19","Version 0.38").replace("Nutzer 08.10.2026","Nutzer 10.10.2026"))
BLD='2026-10-10.5'
ws=wb['Deckblatt']; ws['B4']='0.38'; ws['B6']=ws['B6'].value.replace('Build 2026-10-10.4',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def app(key,txt,src=None):
    r=rows[idx(rows,key)]; r[3]=(r[3] or '')+' '+txt; r[7]=(r[7] or '')+'; '+N+(': '+src if src else '')
app('STE-20','Entscheidung 10.10.2026: Das Tau wird ohne Rückfrage ausgeworfen.','keine Abfrage')
app('KRT-05','Ein Schalter „Lufträume“ rechts unter den Anzeigen schaltet die Ebene während der Fahrt ein und aus; die Einstellung folgt ihm.','Schalter ohne Umweg über die Einstellungen')
app('SPI-18','Im Schritt „Landen“ wird „Weiter“ nach 3 Sekunden klickbar und blendet das Fenster aus; nach der Landung erscheint der letzte Schritt. Nach einer Pause (zum Beispiel Handy ausgeschaltet) setzt „Weiter“ die Fahrt fort. Der Schritt zu den Einstellungen zeigt auch den Schalter für die Lufträume.','Funde vom Handy')
app('SPI-13','Ein Tipp auf die Anzeige „Pause“ setzt die Fahrt fort.','Fund: Pause im Tutorial nicht aufhebbar')
i=max(k for k,r in enumerate(rows) if str(r[0]).startswith('SPI-'))+1
rows.insert(i,['SPI-19','Spielablauf','Weg nach der Landung ansehen','Die Auswertung hat den Eintrag „Weg auf der Karte“. Er schiebt das Ergebnis klein nach links unten und passt die Karte auf den gefahrenen Weg ein. Das kleine Fenster bietet „Auswertung“ (zurück zur vollen Ansicht), „Rückblick“ und „Hauptmenü“.','Nach der Landung ist der ganze Weg auf der Karte zu sehen; Rückblick und Hauptmenü bleiben erreichbar.','Soll','Umgesetzt',N,None,'Basis'])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-10.4',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-40')]; r[2]=r[2]+', SPI-19'; r[5]=r[5]+'; den Schalter „Lufträume“ unter den Anzeigen benutzen; nach der Landung „Weg auf der Karte“ wählen'; r[6]=r[6].replace('Tau auswerfen (Rückfrage)','Tau auswerfen')+'; der Schalter blendet die Lufträume ein und aus; nach der Landung zeigt die Karte den ganzen Weg, das Ergebnis liegt klein links unten'; r[5]=r[5].replace('Tau auswerfen (Rückfrage), noch einmal tippen','Tau auswerfen, noch einmal tippen')
r[11]=(r[11] or '')+f' Build {BLD}: Schalter, Tau ohne Rückfrage und kleines Ergebnisfenster im Browser geprüft.'
r=rows[idx(rows,'T-41')]; r[11]=(r[11] or '')+f' Build {BLD}: Pause im Tutorial mit „Weiter“ und per Tipp auf „Pause“ aufgehoben; Schritt „Landen“ nach 3 s ausblendbar, letzter Schritt erscheint nach der Landung.'
put(ws,5,rows,15,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.38','10.10.2026',f'Build {BLD}: Funde vom Handy. Tutorial: Fenster im Schritt „Landen“ ausblendbar, „Weiter“ hebt eine Pause auf (SPI-18); Tipp auf „Pause“ setzt fort (SPI-13). Nach der Landung „Weg auf der Karte“ mit kleinem Ergebnisfenster (SPI-19). Schalter „Lufträume“ unter den Anzeigen (KRT-05). Schlepptau ohne Rückfrage (STE-20).','Nutzer 10.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
