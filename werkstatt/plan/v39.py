import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.39.xlsx").replace("0.18.xlsx","0.38.xlsx").replace("Version 0.19","Version 0.39").replace("Nutzer 08.10.2026","Nutzer 10.10.2026"))
BLD='2026-10-10.6'
ws=wb['Deckblatt']; ws['B4']='0.39'; ws['B6']=ws['B6'].value.replace('Build 2026-10-10.5',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
print([(r[0],r[2]) for r in rows if str(r[0]).startswith('SRV')])
i=max(k for k,r in enumerate(rows) if str(r[0]).startswith('SRV-'))+1
rows.insert(i,['SRV-07','Server','Statistik für alle','In den Einstellungen öffnet der Knopf „Statistik“ (nur mit Dienst auf dem Server) eine Seite für alle Nutzer: Geräte insgesamt und in den letzten 7 Tagen, gewertete Fahrten, Zeit aller zusammen und die Zeit mit der App je Gerät als Wenigspieler (untere 10 %), die meisten (mittlere Hälfte), Median, Durchschnitt und längste Zeit, dazu eine Verteilung in sieben Klassen. Die App meldet die Zeit, in der sie sichtbar geöffnet ist, jede Minute. Der Dienst gibt nur Summen und Kennwerte heraus.','Die Seite zeigt die Kennwerte; ein zweites Gerät erhöht die Zahl der Geräte.','Soll','Umgesetzt',N+': Die Zeiten sind schief verteilt, nicht wie eine Glockenkurve; deshalb Median und mittlere Hälfte statt Mittelwert und Streuung',None,'Feature'])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-10.5',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-35')]; r[2]=r[2]+', SRV-07'; r[5]=(r[5] or '')+'; in den Einstellungen „Statistik“ öffnen; mit einem zweiten Gerät öffnen und erneut ansehen'; r[6]=(r[6] or '')+'; die Statistik zeigt Geräte und Zeiten, das zweite Gerät wird gezählt'
r[11]=(r[11] or '')+f' Build {BLD}: Statistik mit 34 nachgestellten Geräten geprüft (Median 45 min, Durchschnitt 1 h 09 min, längste Zeit 11 h 23 min); die App meldet die Zeit nach 60 s. Service Worker am Testserver geprüft: App-Dateien gespeichert, Start ohne Netz, Aufrufe des Dienstes gehen am Zwischenspeicher vorbei. Offen bleibt der Betrieb auf dem Synology-NAS.'
put(ws,5,rows,15,4)
ws=wb['Schnittstellen']; rows=rows_of(ws,4,8)
r=rows[idx(rows,'IF-09')]; r[7]=(r[7] or '')+f' Build {BLD}: zusätzlich pubstats (öffentliche Kennwerte) und Meldung der Nutzungszeit jede Minute.'
put(ws,4,rows,8,3)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'OP-34')]; r[7]=(r[7] or '')+f' 10.10.2026: Der Nutzer kann jetzt eine Version auf den Server legen. Paket Gasballon_Server_Build_{BLD}.zip mit Bestenliste und Statistik.'
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.39','10.10.2026',f'Build {BLD}: Statistik für alle über die Einstellungen (SRV-07): Geräte, Fahrten, Verteilung der Nutzungszeit; die App meldet die Zeit jede Minute. Seite „Karten und Daten“ nennt, was der Dienst speichert. Service Worker und Start ohne Netz am Testserver geprüft (T-35).','Nutzer 10.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
