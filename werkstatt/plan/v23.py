import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.23.xlsx").replace("0.18.xlsx","0.22.xlsx").replace("Version 0.19","Version 0.23"))
BLD='2026-10-08.13'
ws=wb['Deckblatt']; ws['B4']='0.23'; ws['B6']=ws['B6'].value.replace('Build 2026-10-08.11',f'Build {BLD}')
for c,src in ((4,1),(5,2)): ws.cell(15,c)._style=copy.copy(ws.cell(15,src)._style); ws.cell(20,c)._style=copy.copy(ws.cell(20,src)._style)
ws['D15']='Status Erledigt (am Handy bestätigt)'; ws['E15']='=COUNTIF(Anforderungen!G5:G300,"Erledigt")'
ws.column_dimensions['D'].width=max(ws.column_dimensions['D'].width or 0,34)

ws=wb['Tests']; trows=rows_of(ws,5,14)
for k in ('T-05','T-08'):
    r=trows[idx(trows,k)]; r[10]='Bestanden'; r[11]='Am Handy in der Web-Version (Build 2026-10-08.12) bestanden. Vorher am Rechner: '+(r[11] or ''); r[12]='08.10.2026'; r[13]='Nutzer'
r=trows[idx(trows,'T-27')]; r[5]=r[5]+'; in der Fahrt oben links „Hauptmenü“ antippen, Rückfrage abbrechen, dann bestätigen'; r[6]=r[6]+'; der Knopf „Hauptmenü“ ist außerhalb des Hauptmenüs immer sichtbar und führt in der Luft erst nach Rückfrage zurück'
r[2]='SPI-08, SPI-09, SPI-10, SPI-13, SPI-14'; r[11]=(r[11] or '')+f' Build {BLD}: Knopf „Hauptmenü“ am Rechner geprüft, am Boden ohne, in der Luft mit Rückfrage.'
r=trows[idx(trows,'T-16')]; r[2]=r[2]+', NFA-06'; r[5]=r[5]+'; im Browser und in der installierten App prüfen, ob das Bedienfeld unten ganz zu sehen ist'; r[6]=r[6]+'; Sandsäcke, Sandmulde, Ventil und Tau liegen vollständig über den Systemtasten des Handys'
put(ws,5,trows,14,4)
passed={r[0] for r in trows if r[10]=='Bestanden'}

ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
ids=[r[0] for r in rows]; done=[]
for r in rows:
    ts=[t for t in trows if r[0] in [x.strip() for x in (t[2] or '').split(',')]]
    if ts and all(t[0] in passed for t in ts) and r[6]=='Umgesetzt': r[6]='Erledigt'; done.append(r[0])
print('Erledigt:',done)
i=idx(rows,'SPI-13')+1
rows.insert(i,['SPI-14','Spielablauf','Fahrt abbrechen und zurück ins Hauptmenü','Außerhalb des Hauptmenüs steht oben links ein Knopf „Hauptmenü“ (Haus-Symbol). Er bricht die Vorbereitung oder die laufende Fahrt ab und führt ins Hauptmenü. Ist der Ballon in der Luft, kommt vorher eine Rückfrage, weil die Fahrt verloren geht. Derselbe Weg bleibt in den Einstellungen.','Der Knopf ist in Vorbereitung, Startfenster, Fahrt und Auswertung sichtbar; am Boden führt er sofort ins Hauptmenü, in der Luft erst nach Bestätigung.','Muss','Umgesetzt',N+': der Weg über die Einstellungen wurde nicht gefunden',None])
rows.append(['NFA-06','Nichtfunktional','Bedienfeld über den Systemtasten','Die App füllt genau die sichtbare Fläche des Fensters. Das Bedienfeld unten (Sandsäcke, Sandmulde, Ventil, Tau) liegt nie hinter den Systemtasten des Handys (Menü, Home, Zurück) oder der Adressleiste, im Browser wie in der installierten App.','Auf dem Handy des Nutzers ist das Bedienfeld im Hoch- und Querformat vollständig zu sehen und zu bedienen.','Muss','Teilweise',N+': Fehler in der Web-Version gemeldet; Ursache vermutet, Behebung am Handy noch nicht bestätigt (OP-44)',None])
r=rows[idx(rows,'INS-01')]; r[7]=(r[7] or '')+'; '+N+': App-Symbol ist jetzt der orange Ballon aus der Kopfzeile (vorher Pixel-Ballon)'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA)
ws['A2']=ws['A2'].value.replace('2026-10-08.11',BLD)

ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
rows.append(['OP-44','Installation','Der untere Rand des Handys verdeckte in der Web-Version das Bedienfeld. In Build 2026-10-08.13 zeichnet die App nicht mehr unter die Systemleisten und richtet ihre Höhe nach der sichtbaren Fläche. Ist das Bedienfeld jetzt ganz zu sehen? Trat der Fehler im Browser oder in der installierten App auf, und mit welchem Handy und Browser?','Behoben / tritt weiter auf','Nutzer','8','A',None,None,'Offen'])
put(ws,5,rows,10,4)

ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.23','08.10.2026',f'Build {BLD}: Knopf „Hauptmenü“ oben links zum Abbrechen der Fahrt (SPI-14). Bedienfeld bleibt über den Systemtasten (NFA-06, OP-44; am Handy noch zu bestätigen). App-Symbol ist der orange Ballon (Build 2026-10-08.12). Tests T-05 und T-08 vom Nutzer am Handy bestanden; zugehörige Anforderungen auf „Erledigt“: '+', '.join(done)+'. Neue Kennzahl „Status Erledigt“ auf dem Deckblatt.','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
