import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.20.xlsx").replace("0.18.xlsx","0.19.xlsx").replace("Version 0.19","Version 0.20"))
ws=wb['Deckblatt']; ws['B4']='0.20'
ws['B6']=ws['B6'].value.replace('Build 2026-10-08.8','Build 2026-10-08.9').replace('Hauptmenü im Stil der App','Startbild im Stil der 80er in einem Rahmen im Stil der App')
ws=wb['Ziele']; rows=rows_of(ws,5,4)
setc(rows,'Z-08',2,'Das Spiel beginnt mit einem Startfenster im Stil der Heimcomputer der 80er (Ballon von der Seite, wirft Sand ab, steigt und sinkt, endlos). Es sitzt in einem Rahmen im Design des eigentlichen Spiels, damit klar ist, dass es zur App gehört. Das Menü bietet Freie Fahrt, Challenge, Bestenliste, Wettkampf und Meisterschaft.',1); setc(rows,'Z-08',3,N,1)
put(ws,5,rows,4,4)

ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
def a(key,**kw):
    m={'C':2,'D':3,'E':4,'F':5,'G':6,'H':7}; r=rows[idx(rows,key)]
    for k,v in kw.items(): r[m[k]]=v
a('SPI-08',C='Startfenster der 80er im Rahmen der App',D='Vor der Spiellogik erscheint das Startfenster im Stil der Heimcomputer der 80er: ein grob gerastertes, bewegtes Bild mit einem Ballon von der Seite, der unten links am Boden beginnt, Sand abwirft, steigt, Gas ablässt und wieder sinkt. Dieses Startfenster sitzt in einem Rahmen im Design des eigentlichen Spiels (Farben, Schrift, Linien und Rundungen der App); der Rahmen trägt oben den Namen des Spiels und die Build-Nummer. Das Menü im Startfenster erscheint nach einer Sekunde und baut sich ruckelnd auf. Im Menü gibt der Spieler seinen Namen ein und wählt Freie Fahrt, Challenge oder Bestenliste; Wettkampf und Meisterschaft sind sichtbar, aber gesperrt. Untermenü, Bestenliste und alle weiteren Fenster bleiben im Stil der App. Im Querformat entfällt die Kopfzeile des Rahmens, damit das Menü Platz hat.',
  E='Das Startbild der 80er ist ringsum vom Rahmen im Stil der App umgeben, im Hoch- und im Querformat; das Menü erscheint nach einer Sekunde; alle fünf Punkte sind zu sehen oder durch Scrollen erreichbar, Wettkampf und Meisterschaft gesperrt.',H=N+'; vorher (0.19) ganz im Stil der App, davor ganz im Stil der 80er')
a('PUN-01',C='Punkte nach Abstand, Landezone 100 m',D='Die Landezone im Zielfahren ist ein Kreis mit 100 m Radius um das Ziel. 100 Punkte, also die volle Punktzahl, gibt es für eine sanfte Landung in der Landezone. Außerhalb fallen die Punkte linear auf 0 am Punkteradius. Der Punkteradius beträgt rund 20 Prozent der Entfernung vom Start zum Ziel, mindestens 300 m: Bei 2,5 km Entfernung gibt es die ersten Punkte ab 500 m Abstand. Die Landezone ist als grüner Kreis, der Punkteradius als dicker, gelb gestrichelter Kreis mit dunklem Rand um das Ziel auf der Karte zu sehen. Die Beschreibung der Aufgabe nennt die Landezone.',
  E='Eine sanfte Landung bis 100 m vom Ziel ergibt 100 Punkte; in der Mitte zwischen Landezone und Punkteradius rund 50; am Punkteradius und dahinter 0. Der grüne Kreis auf der Karte hat 100 m Radius.',H=N+'; vorher volle Punktzahl nur bis 20 m. Radius oder Durchmesser siehe OP-41')
a('STE-20',D='Rechts neben dem Ventil sitzt ein Knopf für das Schlepptau. Es ist 50 m lang, 15 kg schwer, rund 10 cm dick und aus sehr rauen Kokosfasern. Ausgelegt hängt es unter dem Korb. Was vom Tau auf dem Boden liegt, trägt der Boden: Dieses Gewicht zählt nicht mehr zur Masse des Ballons, er wird leichter, und das fängt das Sinken ab. Am Boden bremst das Tau die Fahrt: Bremskraft = Reibungsbeiwert × Gewicht des aufliegenden Taus, entgegen der Fahrtrichtung. Angenommen ist ein Reibungsbeiwert von 3,8 für das Schleppen durch Felder; reine Gleitreibung von Kokosfaser auf Gras läge bei etwa 0,6 bis 1, das dicke Tau pflügt aber durch den Bewuchs. Die Fahrt über Grund folgt aus dem Kräftegleichgewicht von Reibung, Luftkraft des Windes auf die Hülle und Masse des Ballons: Er bremst über einige Sekunden ab, bleibt hinter dem Wind zurück, steht, wenn der Wind die Reibung nicht überwindet, und holt den Wind wieder ein, wenn das Tau abhebt. Der Knopf zeigt die aufliegende Länge in Metern und die Entlastung in kg. Ein Geräusch begleitet das Schleifen.',
  E='Mit ausgelegtem Tau setzt der Ballon bei gleichem Abstieg deutlich sanfter auf als ohne; mit 40 m Tau am Boden fährt er bei 4 m/s Wind langsamer als 0,5 m/s, mit 20 m rund 1,4 m/s; bei 2 m/s Wind steht er ab 20 m Tau am Boden; das Abbremsen dauert einige Sekunden; der Knopf zeigt aufliegende Länge und Entlastung.',H=N+'; Reibungsbeiwert ist eine Annahme, je nach Gelände siehe OP-28')
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA)
ws['A2']=ws['A2'].value.replace('2026-10-08.8','2026-10-08.9')

ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]
ws.cell(40,1).value='Reibungsbeiwert des Schlepptaus am Boden (Bremskraft je Gewicht des aufliegenden Taus)'; ws.cell(40,4).value='Annahme Claude für raues, 10 cm dickes Kokostau in Feldern; abgestimmt auf die Angabe des Nutzers: 40 m am Boden, 4 m/s Wind, Ballon steht fast. Reine Gleitreibung auf Gras wäre etwa 0,6 bis 1. Je nach Gelände siehe OP-28'
r0=ws.max_row+1
extra=[('Zielfahren: Radius der Landezone (volle Punktzahl)',100,'m','Vorgabe Nutzer 08.10.2026: Kreis mit 100 m; als Radius verstanden (PUN-01, OP-41)'),
 ('Zielfahren: kleinster Punkteradius',300,'m','Vorschlag Claude: dreifacher Radius der Landezone; sonst 20 Prozent der Entfernung (PUN-01)'),
 ('Schlepptau: Bremskraft bei 10 / 20 / 30 / 40 m am Boden','112 / 224 / 336 / 447','N','Gerechnet: Reibungsbeiwert 3,8 × 0,3 kg/m × 9,81 m/s² × Länge (STE-20)'),
 ('Schlepptau: Fahrt über Grund bei 4 m/s Wind und 10 / 20 / 30 / 40 m am Boden','2,2 / 1,4 / 0,8 / 0,4','m/s','Gerechnet im Gleichgewicht mit der Luftkraft auf die Hülle (1000 m³, cw 0,45); bei 2 m/s Wind steht der Ballon ab 20 m (STE-20)')]
for k,v in enumerate(extra):
    for c,x in enumerate(v,1): ws.cell(r0+k,c).value=x; ws.cell(r0+k,c)._style=copy.copy(ds[c-1])
    ws.cell(r0+k,2).number_format='General'

ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-12')]; r[5]='Landungen im Ziel, 90 m daneben, zwischen Landezone und Punkteradius, am Punkteradius und dahinter, jeweils sanft und hart; bei zwei verschieden weit entfernten Zielen wiederholen; zwei Ergebnisse mit gleicher Punktzahl und unterschiedlichem Restballast vergleichen'
r[6]=r[6].replace('Sanfte Landung im Ziel ergibt 100 Punkte;','Sanfte Landung in der Landezone (bis 100 m vom Ziel) ergibt 100 Punkte, knapp außerhalb weniger; der grüne Kreis hat 100 m Radius;')
r=rows[idx(rows,'T-27')]; r[6]=r[6].replace('Hauptmenü im Stil der App mit bewegtem Hintergrund','Startbild im Stil der 80er, ringsum im Rahmen der App, im Hoch- und im Querformat; das Menü baut sich nach einer Sekunde ruckelnd auf')
r=rows[idx(rows,'T-29')]; r[5]=r[5].replace('bis 40 m Tau am Boden liegen;','bis 40 m Tau am Boden liegen, danach mit 20 m; beobachten, wie schnell der Ballon abbremst;'); r[6]=r[6].replace('mit 40 m Tau am Boden bleibt er fast stehen','mit 40 m Tau am Boden bleibt er fast stehen, mit 20 m fährt er merklich langsamer als der Wind; das Abbremsen dauert einige Sekunden; der Knopf zeigt aufliegende Länge und Entlastung; die Bremswirkung wirkt auf Feldern glaubhaft')
put(ws,5,rows,14,4)

ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
def o(key,**kw):
    m={'C':2,'D':3,'F':5,'G':6,'H':7,'I':8,'J':9}; r=rows[idx(rows,key)]
    for k,v in kw.items(): r[m[k]]=v
o('OP-05',H='Entschieden (Nutzer 08.10.2026): 100 Punkte bei sanfter Landung in der Landezone, einem Kreis mit 100 m um das Ziel (PUN-01); Punkteradius rund 20 Prozent der Entfernung; Rückmeldung zur Landung in vier Stufen (PUN-08). Offen bleibt die Höhe des Abzugs je Stufe und Aufgabe')
o('OP-28',C='Schlepptau: Der Reibungsbeiwert 3,8 gilt für alle Flächen und ist eine Annahme für raues Kokostau in Feldern. Passt die Bremswirkung zur Erfahrung (bei 4 m/s Wind: 20 m am Boden ergeben rund 1,4 m/s Fahrt, 40 m rund 0,4 m/s)? Soll der Wert später vom Gelände abhängen (Wiese, Acker, Wald, Ort), und woher kommt dann die Art des Geländes? Lässt sich das Tau in der Fahrt einholen, und hängt es beim Schleppen schräg (dann liegt weniger auf)?',D='Ein Wert für alle Flächen (im Spiel 3,8) / Flächennutzung aus OpenStreetMap mit eigenen Werten je Gelände / Wert in den Einstellungen wählbar',
  H='Bestätigt: 50 m lang, 15 kg schwer, rund 10 cm dick, sehr raue Kokosfasern, sehr hoher Widerstand in Feldern; das aufliegende Tau zählt nicht zur Ballonmasse; Bremskraft über einen Reibungsbeiwert und die aufliegende Länge (Nutzer 08.10.2026, STE-20). Offen: Bestätigung des Werts 3,8 und Abhängigkeit vom Gelände')
o('OP-37',H='Entschieden (Nutzer 08.10.2026): Das Startfenster im Stil der 80er kommt zurück und sitzt in einem Rahmen im Design des Spiels (SPI-08). Die Musik im Klang der Heimcomputer bleibt',I='08.10.2026',J='Geschlossen')
o('OP-39',H='Beantwortet (Nutzer 08.10.2026): Physik des Schlepptaus mit Reibungsbeiwert, aufliegender Länge und Entlastung des Ballons (STE-20); Rest in OP-28',I='08.10.2026',J='Geschlossen')
rows.append(['OP-41','Punkte','Landezone im Zielfahren: „Kreis mit 100 m“ ist als Radius 100 m umgesetzt (Durchmesser 200 m). Ist der Radius gemeint oder der Durchmesser? Und wie sollen die Punkte außerhalb der Landezone fallen: wie bisher linear bis zum Punkteradius (20 Prozent der Entfernung, mindestens 300 m) oder gibt es außerhalb keine Punkte?','Radius 100 m / Durchmesser 100 m; außerhalb linear fallend / außerhalb 0','Nutzer','5','A',None,None,'Offen'])
put(ws,5,rows,10,4)

ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.20','08.10.2026','Build 2026-10-08.9: Zielfahren mit Landezone von 100 m Radius für die volle Punktzahl (PUN-01, OP-41); Startfenster im Stil der 80er zurück, in einem Rahmen im Design des Spiels (SPI-08, Z-08, OP-37 entschieden); Schlepptau mit ausgewiesenem Reibungsbeiwert, Kräftebilanz mit Masse des Ballons (Abbremsen, Stehenbleiben, Aufholen) und Anzeige der aufliegenden Länge (STE-20, OP-28, OP-39 beantwortet); Parameter und Tests T-12, T-27, T-29 angepasst.','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
