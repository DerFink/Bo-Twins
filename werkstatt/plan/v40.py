import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.40.xlsx").replace("0.18.xlsx","0.39.xlsx").replace("Version 0.19","Version 0.40").replace("Nutzer 08.10.2026","Nutzer 10.10.2026"))
BLD='2026-10-10.7'
ws=wb['Deckblatt']; ws['B4']='0.40'; ws['B6']=ws['B6'].value.replace('Build 2026-10-10.6',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
i=max(k for k,r in enumerate(rows) if str(r[0]).startswith('SRV-'))+1
rows.insert(i,['SRV-08','Server','Nachrichten an den Admin','Ist in der Datei config.php des Dienstes ein Empfänger eingetragen, schickt der Dienst eine E-Mail, wenn eine Liste zu 80 % oder ganz gefüllt ist (Gordon Bennett 400 und 500, Challenges 80 und 100 Einträge), bei 10, 50, 100, 500 und 1000 Geräten, bei auffälligen Meldungen eines Geräts (sehr viele je Stunde, fünfmal 100 Punkte am Tag, mehr als 2000 km) und wenn die Datenbank nicht schreibbar ist. Jeder Anlass meldet sich einmal. Versand über ein Postfach (SMTP) oder den Mailversand des Servers; ein Aufruf mit dem Admin-Schlüssel schickt eine Probenachricht und nennt Fehler. Empfänger und Zugangsdaten stehen nur auf dem Server.','Die Probenachricht kommt an; ein Anlass erzeugt genau eine Nachricht.','Kann','Umgesetzt',N,None,'Feature'])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-10.6',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-35')]; r[2]=r[2]+', SRV-08'; r[5]=(r[5] or '')+'; nach INSTALLATION_Synology.txt die Teile A, B und C durchgehen; Probenachricht abrufen'; r[6]=(r[6] or '')+'; die Probenachricht kommt im Postfach an'
r[11]=(r[11] or '')+f' Build {BLD}: Nachrichten gegen einen nachgestellten Mailserver geprüft (Probenachricht, 10 Geräte, fünfmal 100 Punkte, 2600 km; je eine Nachricht). Nicht geprüft: verschlüsselter Versand (tls, ssl) an einen echten Anbieter und der Mailversand des NAS.'
put(ws,5,rows,15,4)
ws=wb['Risiken']; rr=ws.max_row; r=rr+1
for c,v in enumerate(['R-23','Technik','Die Nachrichten an den Admin kommen nicht an','Der verschlüsselte Versand ließ sich nur gegen einen Testserver ohne Verschlüsselung prüfen; Anbieter verlangen oft App-Kennwörter','Der Admin erfährt nichts von vollen Listen oder Auffälligkeiten; Spiel und Bestenliste laufen weiter',3,1,None,None,'Probenachricht bei der Installation abrufen (Teil C3 der Anleitung)','Fehlertext der Probenachricht auswerten; Zugangsdaten oder Port anpassen','T-35','Nutzer, Claude','Offen','SRV-08'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(rr,c)._style)
ws.cell(r,8).value=f'=F{r}*G{r}'; ws.cell(r,9).value=f'=IF(H{r}>=15,"hoch",IF(H{r}>=8,"mittel","niedrig"))'
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.40','10.10.2026',f'Build {BLD}: Nachrichten an den Admin per E-Mail (SRV-08, R-23). Neue Anleitung INSTALLATION_Synology.txt mit den Teilen App, Bestenliste und Nachrichten (T-35). Die App selbst ist gegenüber Build 2026-10-10.6 unverändert.','Nutzer 10.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
