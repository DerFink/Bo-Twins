import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.29.xlsx").replace("0.18.xlsx","0.28.xlsx").replace("Version 0.19","Version 0.29").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.5'
ws=wb['Deckblatt']; ws['B4']='0.29'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.4',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
r=rows[idx(rows,'AUF-03')]; r[3]=r[3]+' Solange ein Säckchen fällt, ist ein Pfeifton zu hören, der mit der Resthöhe tiefer wird; beim Aufschlag folgt ein dumpfer Schlag.'; r[7]=(r[7] or '')+'; '+N+': Fallgeräusch'
i=idx(rows,'MEI-11')+1
rows.insert(i,['MEI-12','Meisterschaft','Wetterpaket erstellen','In den Einstellungen lädt ein Knopf das Wetter des Renntages (29. bis 31.08.2026, stündlich) für ein Raster von 374 Orten über Nord- und Osteuropa (47° bis 63° Nord, 2° bis 33,5° Ost) aus dem Archiv der Vorhersagen und speichert es als eine Datei gb69-wetter. Sie enthält je Ort und Stunde Höhe, Temperatur, Wind und Bedeckung auf 13 Flächen bis 300 hPa sowie Bodendruck, Niederschlag und Gesamtbedeckung. Der Abruf läuft in 19 Teilen mit Pausen, zeigt den Fortschritt, lässt sich abbrechen und wartet von selbst, wenn der Dienst bremst. Aus der Datei werden die Vorschau-Routen gerechnet (MEI-11); später kann sie allen Spielern als Wetter dienen.','Nach einigen Minuten liegt die Datei bei den Downloads; sie enthält 374 × 72 Datensätze.','Soll','Teilweise',N+'; mit nachgestelltem Archiv geprüft, am echten Dienst offen',None])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.4',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-37')]; r[2]=r[2]+', MEI-12'; r[5]=r[5]+'; in den Einstellungen das Wetterpaket erstellen'; r[6]=r[6]+'; das Wetterpaket wird in einigen Minuten geladen und als Datei gespeichert'
r[11]=(r[11] or '')+f' Build {BLD}: Wetterpaket mit nachgestelltem Archiv erzeugt und ausgelesen: 374 Orte, 72 Stunden, 1,83 Mio. Werte, Stichproben stimmen; eine Drosselung des Dienstes wurde abgewartet.'
r=rows[idx(rows,'T-31')]; r[6]=r[6]+'; während des Falls ist ein Pfeifton zu hören, beim Aufschlag ein Schlag'
put(ws,5,rows,14,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'OP-48')]; r[7]='Entschieden (Nutzer 09.10.2026): Funktion „Wetterpaket erstellen“ in der App (MEI-12, Build 2026-10-09.5). Nächster Schritt: Nutzer erzeugt die Datei und schickt sie an Claude; daraus entstehen die drei Vorschau-Routen'; r[9]='In Klärung'
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.29','09.10.2026',f'Build {BLD}: Fallgeräusch für das Säckchen (AUF-03). Funktion „Wetterpaket für die Meisterschaft erstellen“ in den Einstellungen (MEI-12, OP-48).','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
