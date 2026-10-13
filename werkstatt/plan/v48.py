import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.48.xlsx").replace("0.18.xlsx","0.47.xlsx").replace("Version 0.19","Version 0.48").replace("Nutzer 08.10.2026","Nutzer 12.10.2026"))
BLD='2026-10-12.1'
ws=wb['Deckblatt']; ws['B4']='0.48'; ws['B6']=ws['B6'].value.replace('Build 2026-10-11.5',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'STE-05')]; r[3]=(r[3] or '')+' Seit Build '+BLD+': Eine gelb gestrichelte Linie „prall“ im Vertikalprofil zeigt die Höhe, in der das jetzige Gas die Hülle ganz füllt (Gastemperatur mit dem aktuellen Abstand zur Luft); entweicht Gas, wandert sie nach oben. Beim Erreichen erscheint einmal „Prallhöhe erreicht“ im Fenster der Crew-Meldung, wieder erst nach einem Füllgrad unter 98 %.'; r[7]=(r[7] or '')+'; '+N+': Wunsch Nutzer 12.10.2026, zwei Entwürfe vorab gezeigt, ohne Höhenangabe und ohne Kennzeichnung am Füllgrad'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-11.5',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.48','12.10.2026',f'Build {BLD}: Prallhöhe als Linie im Vertikalprofil und Meldung beim Erreichen (STE-05).','Nutzer 12.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
