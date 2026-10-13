import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.45.xlsx").replace("0.18.xlsx","0.44.xlsx").replace("Version 0.19","Version 0.45").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
BLD='2026-10-11.3'
ws=wb['Deckblatt']; ws['B4']='0.45'; ws['B6']=ws['B6'].value.replace('Build 2026-10-11.2',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def st(key,status=None,txt=None,note=None):
    r=rows[idx(rows,key)]
    if status: r[6]=status
    if txt: r[3]=(r[3] or '')+' '+txt
    r[7]=(r[7] or '')+'; '+N+(': '+note if note else '')
st('STE-29',None,'Seit Build '+BLD+' ein stetiger Übergang: Die Gasblase oben wird mit sinkendem Füllgrad kleiner, darunter hängt der Stoff als Schlauch; unter 45 % beginnt die Hülle zu flattern, bei 0 % bleibt nur der Schlauch. Kartensymbol und Seitenansicht nutzen dieselbe Formfunktion (hullProf).','Nutzer: Der Wechsel zur Wurst kam zu plötzlich, auch die Ränder schmälern')
st('MEI-11',None,'Seit Build '+BLD+' zeigt der Vorflug „KI-optimiert“ eine Route entlang der Spur des Siegers: tief über die Ostsee, über Kurland und Estland nach Finnland, 1733 km in 56,5 h, Landung bei 62,7° N 27,3° O mit 42 kg Sand.','11.10.2026: Das vorhandene Wetterpaket enthält die Siegerspur bereits (Wind in 100 bis 300 m über Grund). Gefehlt hatten der Suche drei Dinge: Höhenbänder unter 150 m, eine Bewertung entlang der Spur statt nach Strecke je Stunde, und Land in Südfinnland (das Raster von 1° mal 1,5° hält den Streifen 60 bis 61° N für See). Ein feineres Paket war dafür nicht nötig')
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-11.2',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.45','11.10.2026',f'Build {BLD}: Hüllenform geht stetig von der vollen Kugel in die schlaffe, flatternde Hülle über (STE-29). Vorflug „KI-optimiert“ mit neuer Route über die Ostsee nach Finnland, 1733 km (MEI-11).','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
