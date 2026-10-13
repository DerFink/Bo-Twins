import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.35.xlsx").replace("0.18.xlsx","0.34.xlsx").replace("Version 0.19","Version 0.35").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.11'
ws=wb['Deckblatt']; ws['B4']='0.35'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.10',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
r=rows[idx(rows,'MEI-03')]; r[3]=(r[3] or '')+' Länder, die über den 180. Längengrad reichen (Russland), werden fortlaufend gezeichnet; es gibt keine Linien und Sperrstreifen rund um die Erde mehr. Der rote Rand reicht jetzt bis 83° Nord (Nowaja Semlja).'; r[7]=(r[7] or '')+'; '+N+': Fund „zwei Bänder um die ganze Welt im Norden Finnlands, Nowaja Semlja nicht eingetragen“'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.10',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-37')]; r[5]=r[5]+'; die Karte nach Nordfinnland und zur Nowaja Semlja schieben'; r[6]=r[6]+'; in Nordfinnland liegen keine roten Bänder quer über der Karte, Nowaja Semlja ist rot'
r[11]=(r[11] or '')+f' Build {BLD}: Korrektur für den 180. Längengrad mit einem nachgestellten Ring geprüft (größter Sprung danach 80° statt 360°). Mit den echten Grenzdaten am Handy zu prüfen; die Ursache für die fehlende Nowaja Semlja ist nur vermutet.'
r=rows[idx(rows,'T-35')]; r[11]=(r[11] or '')+f' Build {BLD}: Dienst und App erneut geprüft, Ergebnis unverändert.'
put(ws,5,rows,14,4)
ws=wb['Risiken']; 
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.35','09.10.2026',f'Build {BLD}: Sperrzonen über den 180. Längengrad korrigiert; damit entfallen die zwei Bänder quer durch Nordfinnland, die auch disqualifiziert hätten. Roter Rand bis 83° Nord (MEI-03, T-37).','Nutzer 09.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
