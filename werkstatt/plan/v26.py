import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.26.xlsx").replace("0.18.xlsx","0.25.xlsx").replace("Version 0.19","Version 0.26").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.2'
ws=wb['Deckblatt']; ws['B4']='0.26'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.1',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
i=idx(rows,'INS-04')+1
rows.insert(i,['INS-05','Installation','Fenster „Neue Version“','Die App prüft, ob auf dem Server eine neuere Fassung liegt: kurz nach dem Start, wenn sie wieder in den Vordergrund kommt, und alle 30 Minuten. Dazu liest sie die Build-Nummer der Datei auf dem Server; eine eigene Versionsdatei ist nicht nötig. Ist die Nummer höher als die laufende, erscheint ein Fenster mit beiden Nummern und den Knöpfen „Version aktualisieren“ und „Später“. „Version aktualisieren“ lädt die neue Fassung sofort. Während einer Fahrt erscheint das Fenster nicht, sondern danach. Nach „Später“ meldet es sich beim nächsten Start wieder.','Liegt eine höhere Build-Nummer auf dem Server, erscheint das Fenster spätestens beim nächsten Wechsel in die App; nach „Version aktualisieren“ zeigt die App die neue Build-Nummer; bei gleicher oder älterer Nummer erscheint nichts.','Muss','Umgesetzt',N,None])
last=max(int(r[0][4:]) for r in rows if r[0].startswith('REC-')); j=max(k for k,r in enumerate(rows) if r[0].startswith('REC-'))+1
rows.insert(j,[f'REC-{last+1:02d}','Recht','Hinweis: nicht kommerziell, privater Gebrauch','In den Einstellungen steht unter den Hinweisen in allen sieben Sprachen: Die App ist nicht kommerziell und nur für den privaten Gebrauch bestimmt; Nutzung auf eigene Verantwortung, ohne Gewähr.','Der Hinweis ist in jeder Sprache zu lesen.','Muss','Umgesetzt',N+'; der Zusatz „auf eigene Verantwortung, ohne Gewähr“ ist ein Vorschlag von Claude; Bedingungen für eine spätere kommerzielle Nutzung siehe OP-45',None])
print('neu:',rows[j][0])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.1',BLD)
rec=rows[j][0]
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-16')]; r[2]=r[2]+', INS-05'; r[5]=r[5]+'; App offen lassen, neue Version hochladen, App in den Hintergrund und wieder nach vorn holen'; r[6]=r[6]+'; das Fenster „Neue Version“ erscheint, „Version aktualisieren“ lädt die neue Fassung'
r[11]=(r[11] or '')+f' Build {BLD}: Fenster am Rechner mit nachgestellter Server-Antwort geprüft (neuer, gleich, älter, während der Fahrt, „Später“).'
r=rows[idx(rows,'T-19')]; r[2]=r[2]+', '+rec
put(ws,5,rows,14,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
rows.append(['OP-45','Recht','Spätere kommerzielle Nutzung: Die App ist als nicht kommerziell gekennzeichnet. Für einen kommerziellen Betrieb wären vorher zu klären: Wetterdaten (die kostenlose Schnittstelle von Open-Meteo ist für nicht kommerzielle Nutzung gedacht, kommerziell braucht es einen bezahlten Zugang), Kartenkacheln (die öffentlichen Server von OpenStreetMap sind nicht für kommerzielle Last gedacht, ein Kachel-Anbieter wäre nötig), die öffentlichen Dienste Overpass und Nominatim (eigene oder bezahlte Instanz), Impressum, Datenschutzerklärung und Haftung. Die Bedingungen der Dienste sind vor einer Entscheidung im Wortlaut zu prüfen.','Privat lassen / später kommerziell mit bezahlten Diensten','Nutzer','12','C',None,None,'Offen'])
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.26','09.10.2026',f'Build {BLD}: Fenster „Neue Version“ mit Knopf „Version aktualisieren“ (INS-05, T-16). Hinweis „nicht kommerziell, nur privater Gebrauch“ in den Einstellungen ({rec}). Offener Punkt zur späteren kommerziellen Nutzung (OP-45).','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
