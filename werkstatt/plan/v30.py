import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.30.xlsx").replace("0.18.xlsx","0.29.xlsx").replace("Version 0.19","Version 0.30").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.6'
ws=wb['Deckblatt']; ws['B4']='0.30'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.5',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
def app(key,txt,src=None,acc=None):
    r=rows[idx(rows,key)]; r[3]=(r[3] or '')+' '+txt
    if acc: r[4]=acc
    r[7]=(r[7] or '')+'; '+N+(': '+src if src else '')
app('SPI-08','Die Bestenliste steht als letzter Menüpunkt, abgesetzt mit gestricheltem Rahmen und Sternen, mit dem Text „Die besten Fahrten“. Unter dem Namensfeld steht: „Trage hier Deinen Namen ein und finde Dich in der Bestenliste zwischen den anderen Mitspielern.“','Startseite umgestellt')
r=rows[idx(rows,'STE-22')]; r[3]='Am Start hält die Bodencrew den Ballon am Boden fest, bis er sicher steigt: mindestens 30 kg Auftriebsüberschuss (zwei Säcke) und so viel, dass der pralle Ballon bis 250 m über Grund steigt. In einer Bodeninversion, wie sie nachts häufig ist, wird die Luft nach oben wärmer und dünner; dann verlangt die Crew deutlich mehr (gerechnet bei +6 K bis 300 m: 56 kg). Bis dahin steht der Ballon still, und die Fahrzeit läuft nicht. Ein Hinweis nennt den nötigen Überschuss und wie viele kg noch abzuwerfen sind; beim Loslassen kommt eine Meldung. Auch die Vergleichsfahrten starten mit diesem Überschuss.'
r[4]='Nach dem Loslassen steigt der Ballon in jeder Wetterlage auf mindestens rund 250 m; in der Normalatmosphäre reichen zwei Säcke; der Hinweis zählt die fehlenden kg herunter.'; r[7]=(r[7] or '')+'; '+N+': Fund „startet nicht oder erst viele Säcke später“, vermutete Ursache Bodeninversion'
app('STE-08','Die Himmelszeile nennt zusätzlich die Uhrzeit am Ballon auf halbe Stunden (Meisterschaft in MESZ, echtes Wetter in der Zeit des Geräts, Simulation in Sonnenzeit); nachts steht dort nur die Uhrzeit. Eine eigene Zeile zeigt die Lufttemperatur in der Höhe des Ballons.','Uhrzeit und Temperatur')
app('STE-17','Das Profil reicht immer mindestens bis 6000 m; fährt der Ballon höher, wächst die Achse mit.','bis 6000 m')
r=rows[idx(rows,'MEI-09')]; r[2]='Zeitraffer 500× und 1000×'; r[3]='In der Meisterschaft gibt es die Stufen 1×, 10×, 100× und 500× zum Einstellen und zusätzlich 1000× als Verstärker, der nur läuft, solange der Knopf gedrückt ist. Die Stufen 5× und 30× entfallen dort aus Platzgründen.'; r[4]='Mit 500× vergehen rund 500 Sekunden Spielzeit je Sekunde, gedrückt rund 1000; losgelassen gilt wieder die gewählte Stufe.'; r[7]=(r[7] or '')+'; '+N+': 500× fest, 1000× als Verstärker'
i=idx(rows,'SPI-14')+1
rows.insert(i,['SPI-15','Spielablauf','Rückblick auf die Fahrt','Während der Fahrt und nach der Landung lässt sich die zurückliegende Fahrt ansehen. Ein Fahrtschreiber hält alle 10 Sekunden Spielzeit Höhe, Gelände, Steigen, Gastemperatur gegen Luft, Lufttemperatur, Sand, Gas, Wasserlast, Sonne, Wolken, Regen und Auftriebsüberschuss fest (bei langen Fahrten seltener). Der Rückblick zeigt Höhe, Gas gegen Luft und Steigen über der Zeit, für die letzten 30 Minuten, 2 Stunden oder die ganze Fahrt, mit Marken für Ballast und Ventil und Hinterlegung für Nacht und Regen. Durch Antippen wählt man einen Zeitpunkt; dazu stehen die Werte und eine Erklärung, was sich in den 10 Minuten davor geändert hat (Gas kälter oder wärmer mit kg Auftrieb, Wasser, abgegebenes Gas, Sand). Während des Rückblicks ist die Fahrt angehalten. Der Knopf sitzt unten rechts und im Auswertungsfenster.','Nach einem unerwarteten Sinken zeigt der Rückblick den Verlauf und nennt die Ursache in kg; nach dem Schließen läuft die Fahrt weiter.','Muss','Umgesetzt',N+'; nach dem Fortsetzen eines Spielstands beginnt die Aufzeichnung neu',None])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.5',BLD)
ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
for k,v in enumerate([('Start: Mindesthöhe, bis zu der der Ballon nach dem Loslassen steigen muss',250,'m','Vorschlag Claude (STE-22); bestimmt den nötigen Überschuss in einer Inversion'),('Meisterschaft: Verstärker solange gedrückt',1000,'×','Vorgabe Nutzer 09.10.2026; 500× ist feste Stufe (MEI-09)'),('Fahrtschreiber: Abstand der Datensätze',10,'s','Vorschlag Claude; verdoppelt sich ab 3000 Datensätzen (SPI-15)')]):
    for c,x in enumerate(v,1): ws.cell(r0+k,c).value=x; ws.cell(r0+k,c)._style=copy.copy(ds[c-1])
    ws.cell(r0+k,2).number_format='General'
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-22')]; r[6]=r[6]+'; bei einem Start in der Nacht oder am frühen Morgen nennt der Hinweis mehr als 30 kg, und der Ballon steigt danach zügig'
r[11]=(r[11] or '')+f' Build {BLD}: Normalatmosphäre 30 kg, Inversion +6 K bis 300 m 56 kg; danach Gipfel 280 bis 300 m.'
r=rows[idx(rows,'T-08')]; r[6]=r[6]+'; Himmelszeile mit Uhrzeit auf halbe Stunden, eigene Zeile für die Lufttemperatur; Profil bis mindestens 6000 m'; r[10]='Geplant'; r[11]=(r[11] or '')+f' Build {BLD}: Anzeige erweitert (Uhrzeit, Temperatur, Profil bis 6000 m), am Rechner geprüft; am Handy erneut zu bestätigen.'; 
r=rows[idx(rows,'T-27')]; r[2]=r[2]+', SPI-15'; r[5]=r[5]+'; in der Fahrt unten rechts den Rückblick öffnen, einen Zeitpunkt antippen, schließen'; r[6]=r[6]+'; der Rückblick zeigt den Verlauf und nennt, was sich vor dem gewählten Zeitpunkt geändert hat; die Bestenliste steht als letzter Menüpunkt'
r[11]=(r[11] or '')+f' Build {BLD}: Rückblick am Rechner geprüft (52 Datensätze nach 9 min, Zeiger, Erklärung, drei Zeitspannen).'
put(ws,5,rows,14,4)
# STE-08 und STE-17 waren erledigt; nach der Erweiterung wieder „Umgesetzt“, bis T-08 erneut bestanden ist
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
for k in ('STE-08','STE-17'): rows[idx(rows,k)][6]='Umgesetzt'
put(ws,5,rows,9,4,fixA)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
rows+=[['OP-52','Physik','Fund „startet nicht oder erst viele Säcke später“: Im Rechenkern hebt der Ballon mit zwei Säcken immer ab. Vermutete Ursache ist eine Bodeninversion, in der er nach 100 bis 150 m hängen bleibt. Seit Build 2026-10-09.6 verlangt die Crew so viel Überschuss, dass er 250 m erreicht. Trifft das den Fund? In welcher Spielart und zu welcher Uhrzeit trat er auf, und wie hoch kam der Ballon?','Behoben / tritt weiter auf','Nutzer','2','A',None,None,'Offen'],
       ['OP-53','Spielablauf','In der Liste des Nutzers vom 09.10.2026 stand der Punkt „Zeit soll“ ohne weiteren Text. Was war gemeint?','Wunsch nennen / entfällt','Nutzer','7','C',None,None,'Offen']]
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.30','09.10.2026',f'Build {BLD}: Rückblick auf die Fahrt mit Fahrtschreiber und Erklärung (SPI-15). Die Bodencrew lässt erst los, wenn der Ballon 250 m erreicht (STE-22, OP-52). Himmelszeile mit Uhrzeit, Zeile für die Lufttemperatur (STE-08); Vertikalprofil bis mindestens 6000 m (STE-17). Meisterschaft: 500× als feste Stufe, 1000× solange gedrückt (MEI-09). Startseite: Bestenliste unten und abgesetzt, Text „Die besten Fahrten“, Hinweis unter dem Namensfeld (SPI-08). T-08 wieder offen, weil die Anzeige erweitert wurde.','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
