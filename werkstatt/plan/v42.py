import openpyxl, copy
from openpyxl.worksheet.datavalidation import DataValidation
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.42.xlsx").replace("0.18.xlsx","0.41.xlsx").replace("Version 0.19","Version 0.42").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
BLD='2026-10-11.1'
# ---------------- Anforderungen: Stand nachziehen
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def st(key,status=None,scope=None,note=None):
    r=rows[idx(rows,key)]
    if status: r[6]=status
    if scope: r[9]=scope
    if note: r[7]=(r[7] or '')+'; '+N+': '+note
st('KRT-02',scope='Später',note='zurückgestellt, gehört nicht zur Basis 1.0')
st('MET-05',scope='Später',note='zurückgestellt; das Vertikalprofil und der Temp zeigen den Wind')
st('DAT-03',scope='Später',note='zurückgestellt')
st('MET-07','Teilweise','Feature','in der Meisterschaft umgesetzt (Archiv der Vorhersagen, Wetterpaket); zählt zu den Features')
st('REC-02','Umgesetzt',note='Quellen stehen an der Karte und auf der Seite „Karten und Daten“ (DAT-04)')
st('REC-03','Teilweise',note='Zwischenspeicher und Abfragegrenzen sind umgesetzt; die Nutzungsregeln der Dienste sind nicht im Wortlaut geprüft')
st('REC-04','Teilweise',note='Der Standort wird nur für Start, Wetter und Gelände verwendet und nicht an den eigenen Server gemeldet; an den Wetterdienst gehen die Koordinaten aber mit vier Nachkommastellen (rund 10 m), nicht gerundet')
st('NFA-04','Umgesetzt',note='am Testserver geprüft (10.10.2026): App-Dateien gespeichert, Start ohne Netz')
st('NFA-06','Umgesetzt',note='vom Nutzer am 09.10.2026 als in Ordnung gemeldet (OP-44)')
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA)
ws['A2']=ws['A2'].value.replace(' Spalte J trennt den Basisumfang (Fahren in der Freien Fahrt) von den Features (Challenge, Meisterschaft, Server, Spende).',' Spalte J: „Basis“ gehört zur Basis-Version 1.0 (Fahren in der Freien Fahrt), „Feature“ sind Challenge, Meisterschaft, Server und Spende, „Später“ ist zurückgestellt.')
for dv in list(ws.data_validations.dataValidation):
    if str(dv.sqref).startswith('J5'): ws.data_validations.dataValidation.remove(dv)
dv=DataValidation(type='list',formula1='"Basis,Feature,Später"',allow_blank=True); ws.add_data_validation(dv); dv.add('J5:J200')
# ---------------- Tests
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-13')]; r[14]='Später'
r=rows[idx(rows,'T-15')]; r[10]='Am Rechner geprüft'; r[11]=(r[11] or '')+' 10.10.2026: Start ohne Netz am Testserver geprüft (App-Dateien gespeichert, App startet im Offline-Betrieb). Zwischenspeicher des Wetters am Handy zu prüfen.'; r[12]='10.10.2026'; r[13]='Claude (Rechner)'
put(ws,5,rows,15,4); tests=rows
# ---------------- Fahrplan
ws=wb['Fahrplan']; rows=rows_of(ws,5,10)
for k,(s,h) in {'3':('Erledigt','Umgesetzt; Abnahme am Handy in Phase 14'),'4':('Erledigt','Umgesetzt (Feature Challenge)'),'5':('Erledigt','Umgesetzt (Feature Challenge)'),'6':('Geplant','Zurückgestellt, nicht Teil der Basis 1.0'),'7':('In Arbeit','Sprachen und Hinweise umgesetzt; Chaserfahrzeug zurückgestellt'),'8':('In Arbeit','Veröffentlicht auf GitHub und auf dem Server; Abnahme am Handy läuft (Phase 14)'),'9':('Geplant','Zurückgestellt'),'10':('Erledigt','Ersetzt durch Phase 13'),'11':('Erledigt','In Phase 13 aufgegangen'),'12':('In Arbeit','App läuft auf cloud.aerarium.de; Dienst für Bestenliste und Statistik wartet auf PHP am Server'),'13':('In Arbeit','Umgesetzt bis auf offizielle Sperrflächen und Spuren (OP-57)')}.items():
    r=rows[idx(rows,k)]; r[8]=s; r[9]=h
rows.append(['14','Abnahme Basis 1.0','Die Basis-Tests am Handy durchgehen (Blatt Abnahme_Basis) und die Ergebnisse eintragen; verbleibende Basis-Punkte entscheiden.','Alle Basis-Tests bestanden; alle Muss-Anforderungen der Basis umgesetzt; kein offener Punkt mit Prio A, der die Basis betrifft.','Phasen 1 bis 3, 7, 8','Blatt Abnahme_Basis',None,None,'In Arbeit','Die Basis ist gebaut; es fehlt die Bestätigung am Handy'])
put(ws,5,rows,10,4)
# ---------------- Schnittstellen und Lizenzen: Status nachziehen
ws=wb['Schnittstellen']; rows=rows_of(ws,4,8)
S={'IF-01':'In Betrieb','IF-02':'In Betrieb','IF-03':'In Betrieb','IF-04':'In Betrieb','IF-05':'Zurückgestellt','IF-06':'In Betrieb','IF-07':'In Betrieb','IF-08':'In Betrieb','IF-09':'Bereit; wartet auf PHP am Server','IF-10':'In Betrieb','IF-11':'In Betrieb','IF-12':'In Betrieb','IF-13':'In Betrieb'}
for r in rows:
    if r[0] in S: r[6]=S[r[0]]
r=rows[idx(rows,'IF-01')]; r[7]=str(r[7]).replace('Mittelung etwa 15 s','Mittelung 15 s mit Countdown')
put(ws,4,rows,8,3)
ws=wb['Lizenzen']; rows=rows_of(ws,5,6)
for r in rows:
    if r[0] in ('LZ-01','LZ-02','LZ-03','LZ-04','LZ-05','LZ-09'): r[5]='Quellenangabe umgesetzt; Wortlaut nicht rechtlich geprüft'
    if r[0]=='LZ-06': r[5]='Zurückgestellt'
    if r[0]=='LZ-08': r[5]='In Betrieb'
put(ws,5,rows,6,4)
# ---------------- Abnahme_Basis: neues Blatt
if 'Abnahme_Basis' in wb.sheetnames: del wb['Abnahme_Basis']
src=wb['Tests']; ab=wb.create_sheet('Abnahme_Basis', wb.sheetnames.index('Tests')+1)
hs=copy.copy(src.cell(4,1)._style); ts=copy.copy(src.cell(5,4)._style); t1=copy.copy(wb['Tests'].cell(1,1)._style); t2=copy.copy(wb['Tests'].cell(2,1)._style)
ab['A1']='Abnahme der Basis-Version 1.0'; ab['A1']._style=t1
ab['A2']='Die Basis ist gebaut und am Rechner geprüft. „Fertig“ ist sie, wenn (1) alle Basis-Tests am Handy bestanden sind, (2) alle Muss-Anforderungen der Basis umgesetzt sind und (3) kein offener Punkt mit Prio A die Basis betrifft. Trag in Spalte F das Ergebnis ein oder nenn es im Gespräch; Claude überträgt es ins Blatt Tests. Spalte E nennt, was aus deinen bisherigen Rückmeldungen schon dafür spricht; das ersetzt den Test nicht.'; ab['A2']._style=t2
hdr=['Test-ID','Testziel','Stand','Am Handy zu tun (Kurzfassung)','Aus bisherigen Rückmeldungen','Ergebnis am Handy','Bemerkung']
for c,v in enumerate(hdr,1): ab.cell(4,c).value=v; ab.cell(4,c)._style=hs
EV={'T-01':'GPS-Start mit Countdown vom Nutzer benutzt („super gelöst“, 10.10.)','T-02':'Karte wird laufend benutzt; kein offener Fund','T-04':'Zeitraffer laufend benutzt','T-08':'am 08.10. bestanden, seitdem erweitert (Füllgrad, Temperatur, Uhrzeit)','T-14':'echtes Wetter benutzt; kein offener Fund','T-15':'Start ohne Netz am Testserver geprüft','T-16':'App auf dem Server installiert; Fenster „Neue Version“ gesehen','T-18':'','T-19':'Seite „Karten und Daten“ neu seit 10.10.','T-20':'','T-21':'Rückmeldung „Bildschirm zu voll“ in Build 2026-10-11.1 bearbeitet','T-26':'Geräusche in Ordnung; Melodie am 10.10. nicht gehört, seitdem lauter','T-28':'','T-30':'Seitenansicht mit Gelände neu seit 10.10.','T-40':'viele Funde vom 10.10. eingearbeitet','T-41':'Tutorial am Handy durchgespielt; Funde eingearbeitet','T-22':'„viel besser jetzt“ (09.10.)','T-27':'Rückblick benutzt','T-36':'Rückmeldung der Piloten steht aus (OP-46)'}
k=5
for r in tests:
    if r[14] not in ('Basis','Basis und Feature') or r[10]=='Bestanden': continue
    steps=str(r[5] or ''); steps=steps if len(steps)<260 else steps[:257].rsplit(';',1)[0]+' …'
    for c,v in enumerate([r[0],r[3],r[10],steps,EV.get(r[0],''),None,None],1): ab.cell(k,c).value=v; ab.cell(k,c)._style=ts
    k+=1
last=k-1
for col,w in zip('ABCDEFG',(9,30,20,70,48,20,30)): ab.column_dimensions[col].width=w
dv=DataValidation(type='list',formula1='"Bestanden,Nicht bestanden,Nicht geprüft"',allow_blank=True); ab.add_data_validation(dv); dv.add(f'F5:F{last}')
ab.freeze_panes='C5'; ab.auto_filter.ref=f'A4:G{last}'; ab.merge_cells('A2:G2'); ab.row_dimensions[2].height=62
k=last+2; ab.cell(k,1).value='Noch nicht umgesetzte oder nur teilweise umgesetzte Basis-Anforderungen'; ab.cell(k,1)._style=hs
for c in range(2,8): ab.cell(k,c)._style=hs
req=rows_of(wb['Anforderungen'],5,10); PROP={'MET-04':'So lassen: Das Wetter wird je Fahrt geladen; ein Zwischenspeicher über 6 Stunden ist für 1.0 nicht nötig.','REC-03':'Vor einer breiten Weitergabe die Nutzungsregeln von Open-Meteo, OpenStreetMap, Overpass, Nominatim und open flightmaps im Wortlaut lesen.','REC-04':'Koordinaten für den Wetterabruf auf zwei Nachkommastellen runden (rund 1 km); kleine Änderung, auf Zuruf.','NFA-01':'Am Handy beurteilen: Läuft die Karte flüssig?','NFA-02':'Am Handy beurteilen: Sind alle Knöpfe mit einer Hand gut zu treffen?','NFA-03':'Am Handy stoppen: Wie lange dauert „Ziel wird berechnet“?'}
for r in req:
    if r[9]=='Basis' and r[6] in ('Offen','Teilweise'):
        k+=1
        for c,v in enumerate([r[0],r[2],r[6]+' ('+str(r[5])+')',str(r[3])[:300],PROP.get(r[0],''),None,None],1): ab.cell(k,c).value=v; ab.cell(k,c)._style=ts
# ---------------- Deckblatt
ws=wb['Deckblatt']; ws['B4']='0.42'; ws['B5']='11.10.2026'
ws['A2']='Ein Gasballon-Simulator als Web-App für das Handy: Freie Fahrt mit berechnetem Trimmballast, begrenztem Gas, Sonne, Regen und dem Wind der Höhenschichten (Basis), dazu Challenge-Aufgaben, die Nachfahrt des 69. Gordon Bennett und eine gemeinsame Bestenliste (Features).'
ws['B6']=f'Basis-Version 1.0: gebaut und am Rechner geprüft, Abnahme am Handy läuft (Blatt Abnahme_Basis). Build {BLD} ist auf GitHub und auf dem Server veröffentlicht. Features: Challenge umgesetzt; Meisterschaft umgesetzt bis auf offizielle Sperrflächen und Spuren; Bestenliste und Statistik für alle bereit, warten auf PHP am Server. Zurückgestellt: zweiter Ballon, Wettkampf online, Chaserfahrzeug.'
sA=copy.copy(ws['A11']._style); sB=copy.copy(ws['B11']._style); sH=copy.copy(ws['A10']._style)
blk=[('Stand der Basis-Version 1.0',None),('Basis-Anforderungen gesamt','=COUNTIF(Anforderungen!J5:J300,"Basis")'),('davon umgesetzt oder erledigt','=COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!G5:G300,"Umgesetzt")+COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!G5:G300,"Erledigt")'),('davon teilweise','=COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!G5:G300,"Teilweise")'),('davon offen','=COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!G5:G300,"Offen")'),('Umsetzungsgrad der Basis','=E12/E11'),('Muss-Anforderungen der Basis, nicht umgesetzt','=COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!F5:F300,"Muss",Anforderungen!G5:G300,"Offen")+COUNTIFS(Anforderungen!J5:J300,"Basis",Anforderungen!F5:F300,"Muss",Anforderungen!G5:G300,"Teilweise")'),('Basis-Tests gesamt','=COUNTIF(Tests!O5:O300,"Basis")+COUNTIF(Tests!O5:O300,"Basis und Feature")'),('davon am Handy bestanden','=COUNTIFS(Tests!O5:O300,"Basis*",Tests!K5:K300,"Bestanden")'),('davon am Rechner geprüft (Handy offen)','=COUNTIFS(Tests!O5:O300,"Basis*",Tests!K5:K300,"Am Rechner geprüft")'),('davon noch nicht geprüft','=COUNTIFS(Tests!O5:O300,"Basis*",Tests!K5:K300,"Geplant")'),('Zurückgestellt (Umfang „Später“)','=COUNTIF(Anforderungen!J5:J300,"Später")'),('Feature-Anforderungen umgesetzt / gesamt','=COUNTIFS(Anforderungen!J5:J300,"Feature",Anforderungen!G5:G300,"Umgesetzt")&" / "&COUNTIF(Anforderungen!J5:J300,"Feature")')]
for i,(a,b) in enumerate(blk):
    r=10+i; ws.cell(r,4).value=a; ws.cell(r,5).value=b; ws.cell(r,4)._style=sH if i==0 else sA; ws.cell(r,5)._style=sH if i==0 else sB
ws['E15'].number_format='0 %'; ws.column_dimensions['D'].width=46; ws.column_dimensions['E'].width=16
for r in range(26,30):
    v=ws.cell(r,1).value
    if v and str(v).startswith('Blätter:'): ws.cell(r,1).value='Blätter: Ziele, Fahrplan, Anforderungen, Parameter, Schnittstellen, Risiken, Tests, Abnahme_Basis, Offene_Punkte, Lizenzen, Glossar, Änderungen.'
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.42','11.10.2026','Plan auf den Stand „Basis-Version 1.0 in der Abnahme“ gebracht: Deckblatt mit Kennzahlen der Basis, neues Blatt Abnahme_Basis, Fahrplan mit Phase 14 und aktuellen Status, Umfang „Später“ für Zurückgestelltes (KRT-02, MET-05, DAT-03), Status nachgezogen (REC-02, NFA-04, NFA-06 umgesetzt; REC-03, REC-04 teilweise), Schnittstellen und Lizenzen mit Betriebsstatus. Kein neuer Build.','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT,'Abnahmezeilen',last-4)
