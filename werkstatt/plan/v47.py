import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.47.xlsx").replace("0.18.xlsx","0.46.xlsx").replace("Version 0.19","Version 0.47").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
BLD='2026-10-11.5'
ws=wb['Deckblatt']; ws['B4']='0.47'; ws['B6']=ws['B6'].value.replace('Build 2026-10-11.4',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'STE-29')]; r[3]=(r[3] or '')+' Seit Build '+BLD+' gilt die vom Nutzer gewählte Form: bis 50 % Füllgrad Pilzform (oben volle Breite, unten vom Netz schmal gezogen), darunter wird die Hülle auch oben schmaler, zuerst langsam; unter 30 % beginnt sie zu flattern, unter 20 % hängt sie als Wurst am Korb. So lässt sich ein Füllgrad von 60 bis 40 % am Symbol ablesen.'; r[7]=(r[7] or '')+'; '+N+': Auswahl aus drei vorab gezeigten Verläufen (dritte Reihe); Begründung Nutzer: Die Leinen verformen den Ballon in der Realität pilzförmig'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-11.4',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.47','11.10.2026',f'Build {BLD}: Hüllenform nach Wahl des Nutzers, Pilzform bis 50 %, dann schmaler bis zur flatternden Hülle (STE-29). Arbeitsweise ab jetzt: erst prüfen und fragen, bei Darstellungen vorab Bilder, erst nach Freigabe einbauen.','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
