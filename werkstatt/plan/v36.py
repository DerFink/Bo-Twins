import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.36.xlsx").replace("0.18.xlsx","0.35.xlsx").replace("Version 0.19","Version 0.36").replace("Nutzer 08.10.2026","Nutzer 10.10.2026"))
BLD='2026-10-10.3'
ws=wb['Deckblatt']; ws['B4']='0.36'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.11',f'Build {BLD}')
# ---------------- Anforderungen
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
def app(key,txt,src=None,st=None):
    r=rows[idx(rows,key)]; r[3]=(r[3] or '')+' '+txt
    if st: r[6]=st
    r[7]=(r[7] or '')+'; '+N+(': '+src if src else '')
app('SPI-08','Entscheidung 10.10.2026: Das Startbild im Stil der 80er entfällt, weil es von der App ablenkt und irreführt; an seine Stelle tritt die Startseite mit Seitenansicht (SPI-17).','Retro-Bild entfällt')
app('STE-01','Wortwahl seit Build '+BLD+': Sandsäcke werden entleert und Sand wird abgegeben, nicht abgeworfen; geworfen werden nur die Säckchen im Wandziel.','Wortwahl')
app('STE-08','Statt „Gas in Prozent zum Start“ zeigt die Anzeige den Füllgrad der Hülle in Prozent (STE-25); der Wert zum Start steht im Detailfenster (STE-26) und im Rückblick.','Füllgrad statt Gas zum Start','Umgesetzt')
app('STE-20','Das Tau lässt sich in einer Fahrt nur einmal auswerfen und nicht wieder einrollen; vor dem Auswerfen fragt die App nach. Der Rückblick markiert den Zeitpunkt mit einer gestrichelten Linie.','einmalig, mit Rückfrage','Umgesetzt')
app('SPI-15','Der Rückblick nennt Füllgrad und Gas zum Start und markiert das Auswerfen des Schlepptaus.','Tau im Rückblick','Umgesetzt')
app('MEI-15','Die Zeile heißt jetzt „Platz“ und zeigt zum Beispiel 14/24; vorher lief sie über den Rand.','Fund am Rechner')
new={'STE':[['STE-25','Steuerung und Physik','Füllgrad bestimmt Form und Widerstand','Der Füllgrad ist das Gasvolumen bei Druck und Gastemperatur am Ort, geteilt durch das Hüllenvolumen. Er bestimmt die Widerstandsfläche (cw × Fläche): prall 54 m² wie bisher, leer 6 m². Dazwischen folgt sie einer Exponentialkurve, die anfangs sehr wenig und am Ende sehr deutlich wirkt: 70 % 53 m², 50 % 51 m², 20 % 37 m², 10 % 25 m². Der leere Ballon fällt mit rund 50 m/s.','Im Rechenkern fällt der Ballon ohne Gas mit 50 bis 55 m/s; bis 50 % Füllgrad ändert sich das Sinken um höchstens wenige Prozent.','Muss','Umgesetzt',N+': Angabe eines Ballonpiloten (50 m/s statt 15 m/s) und Vorgabe „exponentiell“; die Kurve selbst ist eine Annahme',None],
 ['STE-26','Steuerung und Physik','Ballon im Detail','Ein Tipp auf den Ballon in der Karte öffnet ein Fenster mit Füllgrad, Gas zum Start, Gasvolumen, Gastemperatur, Lufttemperatur, Differenz, Auftriebsüberschuss, Gesamtmasse, Sand, Masse des Gases, Wasser auf der Hülle, Schlepptau und Widerstandsfläche. Die Werte laufen mit; ein Tipp schließt.','Das Fenster öffnet sich und zeigt beide Prozentwerte.','Soll','Umgesetzt',N,None],
 ['STE-27','Steuerung und Physik','Temp als Stüve-Diagramm','Ein Tipp auf die linke Hälfte des Vertikalprofils öffnet das Temp-Diagramm, ein Tipp darauf schließt es wieder (ohne Knopf). Stüve-Diagramm: Temperatur waagerecht, Druck senkrecht, Trockenadiabaten als Geraden, Temperaturkurve in den Farben der Schichtung, Windfahnen auf den Druckflächen, Höhen in m, Wolken, Marke für Ballon und Gastemperatur. Ohne Taupunkt und Feuchtadiabaten, weil keine Feuchte geladen wird. Die rechte Hälfte des Profils vergrößert es wie bisher.','Diagramm öffnet und schließt per Tipp; die Kurve passt zum Vertikalprofil.','Soll','Umgesetzt',N+': Bauart Stüve von Claude gewählt, Nutzer sieht es sich an',None],
 ['STE-28','Steuerung und Physik','Seitenansicht mit Gelände','Unter 50 m über Grund erscheint von selbst eine Seitenansicht: Ballon (Form nach Füllgrad), Korb, Schlepptau (hängend und aufliegend), Gelände als Schnitt in Fahrtrichtung, Höhe über Grund und Maßstab. Sand rieselt, wenn Ballast abgegeben wird; am Ventil tritt Gas aus. Ein Tipp blendet sie aus, ein Knopf blendet sie in jeder Höhe ein; in den Einstellungen lässt sich die Automatik abschalten.','Am Start und unter 50 m sichtbar, ausblendbar und in der Höhe einblendbar; ansteigendes Gelände voraus ist zu sehen.','Soll','Umgesetzt',N+': Wunsch eines Ballonpiloten; Gelände aus dem geladenen Raster (400 m), ohne Bäume und Gebäude',None]],
'KRT':[['KRT-05','Karte','Lufträume ein- und ausblendbar','In den Einstellungen lässt sich die Luftfahrtkarte von open flightmaps als Ebene über die Karte legen, wie in der Tracker-App. Sichtbar ab Zoomstufe 7, auf dem Gerät höchstens 14 Tage gespeichert, nur Anzeige ohne Wirkung auf das Spiel. Nicht alle Länder sind abgedeckt (IF-13).','Ebene lässt sich ein- und ausschalten; in Deutschland sind Lufträume zu sehen.','Soll','Umgesetzt',N,None]],
'SPI':[['SPI-17','Spielablauf','Startseite mit Seitenansicht','Die Startseite zeigt oben den Ballon von der Seite über einer vorbeiziehenden Landschaft mit Höhenskala und Windfahnen: Er steigt, während Sand rieselt, und sinkt, während am Ventil Gas austritt; tief über dem Boden schleift das Tau. Darunter stehen Name, Erste Fahrt, Freie Fahrt, Challenge, Meisterschaft, Die besten Fahrten und Buy me a Beer im Stil der App. Bei „Bewegung reduzieren“ steht das Bild.','Die Seite erklärt die App auf einen Blick und wirkt nicht wie ein Retro-Spiel.','Soll','Umgesetzt',N+': Vorschlag B gewählt',None],
 ['SPI-18','Spielablauf','Tutorial „Erste Fahrt“','Ein eigener Menüpunkt führt in 13 Schritten durch eine Fahrt mit simuliertem Wetter ab Bitterfeld. Am Anfang ist nur der Ballon zu sehen; Sandsäcke, Anzeigen, Vertikalprofil, Schaufel, Ventil, Zeitraffer und Schlepptau kommen einzeln dazu. Eine Hand zeigt auf das Element, das anzutippen ist; danach erklärt ein Satz, was geschehen ist und wozu es dient. Bis zur ersten abgeschlossenen Fahrt ist der Menüpunkt hervorgehoben. Texte in Deutsch und Englisch, andere Sprachen zeigen Englisch.','Ein neuer Nutzer hebt ab, trimmt, zieht das Ventil, wirft das Tau aus und landet, ohne die App zu kennen.','Soll','Umgesetzt',N,None]],
'TON':[['TON-02','Ton','Ruhige Melodie auf der Startseite','Die Startseite spielt ein frei erfundenes, ruhiges Stück aus weichen Flächenklängen und einzelnen Glockentönen. Es blendet in 4 Sekunden ein und beim Beginn einer Fahrt in rund 3 Sekunden aus. Die bisherige Melodie im Stil der Heimcomputer entfällt; alle Geräusche der Fahrt und der Jubel am Ziel bleiben.','Die Musik setzt sanft ein und endet sanft; sie lässt sich in den Einstellungen abschalten.','Kann','Umgesetzt',N+': am Rechner nicht anhörbar, Gefallen prüft der Nutzer (OP-60)',None]],
'DAT':[['DAT-04','Daten','Seite „Karten und Daten“','In den Einstellungen steht eine Seite, die für jede Quelle nennt, was geladen wird, wie weit die Abdeckung reicht, wie alt die Daten sein können und unter welchen Bedingungen sie stehen: Karte, Lufträume, Wetter, Wetter der Meisterschaft, Gelände, Zielorte und Ortsnamen, Ländergrenzen. Deutsch und Englisch.','Die Abdeckung der Lufträume und die Grenzen der Daten sind in der App nachlesbar.','Soll','Umgesetzt',N+': „muss dokumentiert und für den Nutzer nachlesbar sein“',None]]}
for pre,lst in new.items():
    last=max(i for i,r in enumerate(rows) if str(r[0]).startswith(pre+'-'))
    rows[last+1:last+1]=lst
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.11',BLD)+' Spalte J trennt den Basisumfang (Fahren in der Freien Fahrt) von den Features (Challenge, Meisterschaft, Server, Spende).'
FEAT=('ZIE','AUF','PUN','KI','SRV','MEI'); FEATID=('SPI-10','SPI-11','SPI-12','SPI-16')
def scope(i): i=str(i); return 'Feature' if i.split('-')[0] in FEAT or i in FEATID else 'Basis'
ws.cell(4,10).value='Umfang'; ws.cell(4,10)._style=copy.copy(ws.cell(4,9)._style); ws.column_dimensions['J'].width=11
for k,r in enumerate(rows): c=ws.cell(5+k,10); c.value=scope(r[0]); c._style=copy.copy(ws.cell(5+k,6)._style)
ws.auto_filter.ref=f'A4:J{4+len(rows)}'
for dv in ws.data_validations.dataValidation:
    s=str(dv.sqref)
    if s.startswith('F5') or s.startswith('G5'): dv.sqref=openpyxl.worksheet.cell_range.MultiCellRange(s[:2]+':'+s[0]+str(4+len(rows)+30))
dv=openpyxl.worksheet.datavalidation.DataValidation(type='list',formula1='"Basis,Feature"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'J5:J{4+len(rows)+30}')
nreq=len(rows)
# ---------------- Parameter
ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
for k,v in enumerate([('Widerstandsfläche der leeren Hülle mit Korb (cw × Fläche)',6,'m²','Aus der Angabe eines Ballonpiloten: rund 50 m/s im freien Fall (STE-25)'),('Form der Kurve Füllgrad zu Widerstand (Exponent)',5,'–','Annahme Claude nach Vorgabe „anfangs gering, am Ende deutlich“ (STE-25, R-21)'),('Seitenansicht: erscheint von selbst unter dieser Höhe über Grund',50,'m','Vorgabe Nutzer 10.10.2026 (STE-28)'),('Lufträume: Zwischenspeicher höchstens',14,'Tage','wie in der Tracker-App (KRT-05)'),('Lufträume: sichtbar ab Zoomstufe',7,'–','Kacheln von open flightmaps gibt es von Stufe 7 bis 12 (KRT-05)'),('Musik der Startseite: Einblenden',4,'s','Vorschlag Claude (TON-02)'),('Tutorial: „Weiter“ nach dem Abheben erst ab Höhe über Grund',30,'m','Vorschlag Claude, damit Ventil und Tau nicht am Boden erklärt werden (SPI-18)')]):
    for c,x in enumerate(v,1): ws.cell(r0+k,c).value=x; ws.cell(r0+k,c)._style=copy.copy(ds[c-1])
    ws.cell(r0+k,2).number_format='General'
# ---------------- Schnittstellen
ws=wb['Schnittstellen']; rows=rows_of(ws,4,8)
def note(k,t): r=rows[idx(rows,k)]; r[7]=(r[7] or '')+' '+t
note('IF-02','Abdeckung: weltweit, stündlich, 13 Druckflächen 1000 bis 250 hPa; Auflösung je nach Region; Feuchte wird nicht geladen. In der App nachlesbar (DAT-04).')
note('IF-03','Abdeckung: weltweit. In der App nachlesbar (DAT-04).')
note('IF-10','Höhenmodell Copernicus, rund 90 m Raster; ohne Bäume und Gebäude. In der App nachlesbar (DAT-04).')
note('IF-11','Archivierte Vorhersage, nicht gemessenes Wetter. In der App nachlesbar (DAT-04).')
note('IF-12','Vereinfachte Grenzen, Abweichung einige Kilometer möglich; Länder über den 180. Längengrad werden fortlaufend gerechnet. In der App nachlesbar (DAT-04).')
rows.append(['IF-13','Spiel','Luftfahrtkarte open flightmaps','HTTPS, Bildkacheln (nwy-tiles-api.prod.newaydata.com, Pfad latest/aero/latest)','Lufträume und Luftfahrtangaben als Bild über der Karte, Zoomstufen 7 bis 12','Server zu Spiel','Offen','Wie in der Tracker-App, ohne Schlüssel. Abdeckung laut openflightmaps.org, Stand 10.10.2026: Belgien, Dänemark, Deutschland, Finnland, Frankreich, Griechenland, Italien, Kroatien, Malta, Namibia, Niederlande, Österreich, Polen, Rumänien, Schweden, Schweiz, Slowakei, Slowenien, Südafrika, Tschechien, Ungarn; nicht abgedeckt unter anderem das Baltikum. Zwischenspeicher 14 Tage. Nur Anzeige (KRT-05, DAT-04).'])
put(ws,4,rows,8,3)
# ---------------- Lizenzen
ws=wb['Lizenzen']; rows=rows_of(ws,5,6)
rows.append(['LZ-09','open flightmaps (Luftfahrtkarte)','Lufträume als Kartenebene (KRT-05)','OFMA General Users’ License: nach Angabe auf openflightmaps.org weltweit, gebührenfrei, nicht ausschließlich, gewerbliche Nutzung eingeschlossen (gelesen am 10.10.2026). Quellenangabe „© open flightmaps association“.','Quellenangabe an der Karte und auf der Seite „Karten und Daten“; Hinweis „nur zur Orientierung“. Lizenztext vor einer breiten Weitergabe im Wortlaut prüfen.','Offen'])
put(ws,5,rows,6,4)
# ---------------- Glossar
ws=wb['Glossar']
for r in range(1,ws.max_row+1):
    if ws.cell(r,1).value=='Ballast': ws.cell(r,2).value='Sand in Säcken und in der Sandmulde; wird er abgegeben (Sack entleeren, Schaufel), steigt der Ballon.'
    if ws.cell(r,1).value=='Sandmulde': ws.cell(r,2).value='Offener Behälter mit losem Sand im Korb, aus dem mit einer kleinen Schaufel jeweils 200 g abgegeben werden.'
r=ws.max_row+1
for k,(a,b) in enumerate([('Füllgrad','Gasvolumen bei Druck und Gastemperatur am Ort, geteilt durch das Hüllenvolumen, in Prozent. Sagt, wie prall die Hülle ist, und bestimmt Form und Luftwiderstand.'),('Gas zum Start','Gasmenge an Bord in Prozent der Menge beim Start. Sagt, wie viel Gas schon verloren ist.'),('Stüve-Diagramm','Meteorologisches Diagramm mit Temperatur waagerecht und Druck senkrecht, in dem Trockenadiabaten Geraden sind.'),('Seitenansicht','Kleines Bild mit Ballon, Schlepptau und Gelände in Fahrtrichtung.')]):
    ws.cell(r+k,1).value=a; ws.cell(r+k,2).value=b
    for c in (1,2): ws.cell(r+k,c)._style=copy.copy(ws.cell(r-1,c)._style)
# ---------------- Risiken
ws=wb['Risiken']; rr=ws.max_row
def risk(i,row):
    r=rr+i
    for c,v in enumerate(row,1): ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(rr,c)._style)
    ws.cell(r,8).value=f'=F{r}*G{r}'; ws.cell(r,9).value=f'=IF(H{r}>=15,"hoch",IF(H{r}>=8,"mittel","niedrig"))'
risk(1,['R-21','Physik','Der Zusammenhang zwischen Füllgrad und Luftwiderstand trifft die Wirklichkeit nicht','Belegt sind nur die beiden Enden (prall und leer); die Kurve dazwischen ist angenommen','Sinken und Steigen bei halb leerem Ballon wirken falsch',3,2,None,None,'Rückmeldung erfahrener Piloten einholen (OP-46)','Exponent und Endwert als Parameter anpassen','T-39','Nutzer, Claude','Offen','STE-25'])
risk(2,['R-22','Technik','Die Luftfahrtkarte fällt aus oder ändert ihre Adresse','Der Kacheldienst von open flightmaps wird ohne Schlüssel und ohne Zusage genutzt','Lufträume werden nicht angezeigt; das Spiel läuft weiter',2,1,None,None,'Ebene ist abschaltbar; Adresse steht an einer Stelle im Code','Adresse anpassen oder Ebene ausblenden','T-40','Claude','Offen','KRT-05, IF-13'])
# ---------------- Tests
ws=wb['Tests']; rows=rows_of(ws,5,14)
rows+=[['T-39','Software','STE-25','Füllgrad und Widerstand','Rechenkern in Node','Widerstandsfläche für 100, 70, 50, 20, 10 und 0 % ausgeben; aus 3000 m das Gas auf 80, 50, 20, 5 und 0 % verringern und fallen lassen; einen normalen Start rechnen','54 / 53 / 51 / 37 / 25 / 6 m²; größtes Sinken ohne Gas 50 bis 55 m/s; der Start bleibt unverändert (zwei Säcke)','Rechner','Jederzeit','2 Physik','Bestanden',f'Build {BLD}: 54,4 / 53,3 / 50,7 / 36,8 / 25,2 / 6,0 m². Größtes Sinken: 80 % Gas 7,7 m/s; 50 % 13,1; 20 % 19,6; 5 % 31,5; 0 % 52,9 m/s. Start mit zwei Säcken, 1,0 m/s nach 150 s.','10.10.2026','Claude (Rechner)'],
 ['T-40','System','STE-08, STE-20, STE-26, STE-27, STE-28, KRT-05, DAT-04, SPI-15','Anzeige, Diagramm, Seitenansicht, Lufträume','Handy, Freie Fahrt mit echtem Wetter in Deutschland','Abheben und auf die Seitenansicht achten; steigen; Seitenansicht über den Knopf ein- und per Tipp ausblenden; links aufs Profil tippen, Diagramm ansehen, antippen; auf den Ballon tippen; Tau auswerfen (Rückfrage), noch einmal tippen; unter 50 m sinken; Rückblick öffnen; in den Einstellungen Lufträume einschalten und „Karten und Daten“ lesen','Die Anzeige zeigt den Füllgrad; Seitenansicht mit Sand, Gas, Tau und Gelände; Stüve-Diagramm öffnet und schließt per Tipp; Detailfenster mit beiden Prozentwerten; das Tau lässt sich nicht einrollen; der Rückblick markiert das Tau; Lufträume erscheinen ab Zoomstufe 7','Handy','Einmal','3 Bedienung','Am Rechner geprüft',f'Build {BLD}: alle Schritte im Browser mit nachgestellter Karte geprüft. Offen am Handy: Tipp auf den echten Ballon in der Karte, Aussehen der Lufträume, Gelände in der Seitenansicht mit echten Höhen, Lesbarkeit des Diagramms.','10.10.2026','Claude (Rechner)'],
 ['T-41','System','SPI-17, SPI-18, TON-02','Startseite, Melodie und Tutorial','Handy, App neu geöffnet','Startseite ansehen und anhören; „Erste Fahrt“ starten und allen 13 Schritten folgen; landen; zurück ins Menü','Das Bild läuft ruhig; die Musik setzt sanft ein und blendet beim Start aus; im Tutorial erscheinen die Bedienelemente einzeln, die Hand zeigt auf das richtige Element; nach der Landung ist der Menüpunkt nicht mehr hervorgehoben','Handy','Einmal','3 Bedienung','Am Rechner geprüft',f'Build {BLD}: Startseite und alle 13 Schritte im Browser durchlaufen; Menü ist auf 740 px Höhe rund 80 px länger als das Fenster und rollt. Die Melodie ließ sich am Rechner nicht anhören.','10.10.2026','Claude (Rechner)']]
put(ws,5,rows,14,4)
import re
ws.cell(4,15).value='Umfang'; ws.cell(4,15)._style=copy.copy(ws.cell(4,14)._style); ws.column_dimensions['O'].width=14
for k,r in enumerate(rows):
    ids=re.findall(r'[A-Z]{2,3}-\d+',str(r[2])); ids=[i for i in ids if not i.startswith(('R-','OP-','T-','IF-','LZ-'))]; s=set(scope(i) for i in ids)
    c=ws.cell(5+k,15); c.value='Basis' if s=={'Basis'} else 'Feature' if s=={'Feature'} else 'Basis und Feature'; c._style=copy.copy(ws.cell(5+k,2)._style)
ws.auto_filter.ref=f'A4:O{4+len(rows)}'
# ---------------- Offene Punkte
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
rows+=[['OP-60','Ton','Gefällt die neue ruhige Melodie der Startseite? Sie ließ sich am Rechner nicht anhören. Lautstärke, Tempo und Klang sind an einer Stelle einstellbar.','So lassen / leiser / langsamer / anderer Klang','Nutzer','3','C',None,None,'Offen'],
 ['OP-61','Physik','Temp-Diagramm: Stüve ohne Taupunkt und Feuchtadiabaten umgesetzt. Passt die Bauart, und soll die Feuchte zusätzlich geladen werden (mehr Daten je Abruf)?','So lassen / Skew-T / Taupunkt dazu','Nutzer','3','C',None,None,'Offen'],
 ['OP-62','Sprachen','Tutorial und die Seite „Karten und Daten“ gibt es nur auf Deutsch und Englisch; die anderen fünf Sprachen zeigen Englisch. Übersetzen?','Übersetzen / so lassen','Nutzer','7','C',None,None,'Offen'],
 ['OP-63','Physik','Fund am Rechner: Hebt der Ballon nur wenige Meter ab (unter 10 m) und sinkt zurück, gilt das nicht als Landung; er steht am Boden, die Zeit bleibt stehen, bis wieder Sand abgegeben wird. Bestand schon vorher. Soll das als Landung zählen?','So lassen / als Landung werten','Nutzer','3','C',None,None,'Offen'],
 ['OP-64','Karte','Lufträume: Die Abdeckung von open flightmaps endet vor dem Baltikum; in der Meisterschaft bleibt die Ebene dort leer. Reicht das, oder soll eine zweite Quelle geprüft werden?','So lassen / zweite Quelle prüfen','Nutzer','7','C',None,None,'Offen']]
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.36','10.10.2026',f'Build {BLD}: Basis und Features getrennt (Spalte „Umfang“ in Anforderungen und Tests). Physik: Füllgrad bestimmt den Widerstand (STE-25, T-39, R-21). Anzeige: Füllgrad, Detailfenster (STE-26), Temp als Stüve-Diagramm (STE-27), Seitenansicht mit Gelände (STE-28). Schlepptau nur einmal und im Rückblick markiert (STE-20). Lufträume von open flightmaps (KRT-05, IF-13, LZ-09, R-22) und Seite „Karten und Daten“ (DAT-04). Neue Startseite (SPI-17), ruhige Melodie (TON-02), Tutorial (SPI-18); Retro-Bild entfällt (SPI-08). Sandsäcke werden entleert (STE-01). Neu: T-40, T-41, OP-60 bis OP-64.','Nutzer 10.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT,nreq,'Anforderungen')
