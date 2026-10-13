import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.24.xlsx").replace("0.18.xlsx","0.23.xlsx").replace("Version 0.19","Version 0.24"))
BLD='2026-10-08.14'
ws=wb['Deckblatt']; ws['B4']='0.24'; ws['B6']=ws['B6'].value.replace('Build 2026-10-08.13',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
r=rows[idx(rows,'AUF-03')]
r[3]='An Bord sind so viele farbige Säckchen wie Target Cross. Ein Säckchen wiegt 70 g. Abgeworfen wird über einen eigenen Knopf rechts unten, der auch die Zahl der übrigen Säckchen zeigt. Das Säckchen fällt in der Spielzeit mit 9 m/s und treibt in jeder Höhe mit dem dortigen Wind, so wie der Ballon in seiner Höhe; während des Falls ist es als kleiner Punkt auf der Karte zu sehen. Erst beim Aufschlag wird der Landepunkt markiert und der Abstand zum Kreuz gemeldet. Bei gleichmäßigem Wind landet es deshalb nahe unter dem Ballon, bei Windscherung daneben. Landet der Ballon, während ein Säckchen noch fällt, zählt es mit. Die Fahrt endet mit der Landung; wo gelandet wird, ist egal.'
r[4]='Der Knopf zeigt die Zahl der Säckchen; jeder Abwurf verringert sie um eins und macht den Ballon um 70 g leichter; das Säckchen ist im Fall zu sehen und schlägt nach Höhe geteilt durch 9 m/s Sekunden auf; bei gleichem Wind in allen Höhen liegt es dann nahe unter dem Ballon.'
r[7]=N+': Fehler behoben (Landepunkt erschien sofort, 100 bis 200 m vor dem Ballon)'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-08.13',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-31')]; r[5]=r[5]+'; beim Abwurf das fallende Säckchen und den Ballon auf der Karte beobachten'; r[6]=r[6]+'; das Säckchen fällt sichtbar und landet bei gleichmäßigem Wind nahe unter dem Ballon'
r[11]=(r[11] or '')+f' Build {BLD}: Abwurf aus 270 m bei 2,7 m/s Wind: Aufschlag nach 30 s, 73 m vom Abwurfort und 8 m vom Ballon; auch mit Zeitraffer und bei Landung während des Falls geprüft.'
put(ws,5,rows,14,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.24','08.10.2026',f'Build {BLD}: Wandziel, Fehler behoben: Das Säckchen fällt jetzt in der Spielzeit und ist dabei zu sehen; vorher erschien der Landepunkt sofort und lag deshalb 100 bis 200 m vor dem Ballon (AUF-03, T-31).','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
