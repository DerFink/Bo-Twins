import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.51.xlsx").replace("0.18.xlsx","0.50.xlsx").replace("Version 0.19","Version 0.51").replace("Nutzer 08.10.2026","Nutzer 26.10.2026"))
BLD='2026-10-26.1'
ws=wb['Deckblatt']; ws['B4']='0.51'; ws['B6']=ws['B6'].value.replace('Build 2026-10-13.1',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def ins(pre,row):
    i=max(k for k,r in enumerate(rows) if str(r[0]).startswith(pre+'-'))+1; rows.insert(i,row)
ins('MEI',['MEI-14','Meisterschaft','Szenarien und 62. Gordon Bennett 2018','Die Meisterschaft ist in Szenarien gegliedert; ein Fenster bietet „69. Gordon Bennett 2026“ (Gladbeck) und „62. Gordon Bennett 2018“ (Bern) mit Hinweis auf gespeicherte Spielstände. Je Szenario: Startplatz, Startzeit, echte Teams mit Landeorten, gesperrte Länder, Wetterquelle, Spielstand und eigene Bestenliste (auf dem Gerät und im Serverdienst, c=gb69/gb62). 2018: 20 Teams aus dem offiziellen Endergebnis; Landepunkte in Richtung des genannten Ortes in der offiziellen Entfernung, „Le Peyrat“ (Rang 11) ohne Punkt, Rang 15 als AUS (Quellen widersprechen sich); Start 28.09.2018 19:00 MESZ auf der Berner Allmend als Annahme; keine gesperrten Länder, Wasserlandung disqualifiziert; Wetter aus der Reanalyse ERA5 (Open-Meteo-Archiv, für Zeitpunkte vor 2022); noch ohne Vorflug.','Beide Szenarien lassen sich wählen, starten, speichern und fortsetzen, ohne sich gegenseitig zu stören; 2018 lädt das Wetter aus dem Archiv; die Listen bleiben getrennt.','Soll','Umgesetzt','Wunsch Nutzer 26.10.2026; Daten: legends.gordonbennett.aero (Ergebnis 62. Rennen)',None,'Feature'])
ins('MEI',['MEI-15','Meisterschaft','Eigenes Team: Land frei wählbar','Statt „JAM-1 CoolRidings“ startet der Spieler unter seinem Namen. Im Auswahlfenster lässt sich ein Land wählen (76 Länder, Namen in der Sprache der App, mit Flagge) oder „ohne Land“. Mit Land fährt er als nächstes Team dieses Landes (z. B. GER-4); das Kürzel steht mit Flagge am Ballon, in Rangliste und Bestenliste.','Mit Land erscheint das Kürzel an Ballon und Ranglisten, ohne Land nur der Name.','Soll','Umgesetzt','Wunsch Nutzer 26.10.2026 (ersetzt die Vorgabe „nur Name“ vom 09.10.)',None,'Feature'])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-13.1',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.51','26.10.2026',f'Build {BLD}: Meisterschaft mit Szenarien, neu 62. Gordon Bennett 2018 Bern (MEI-14); Land frei wählbar statt JAM-1 (MEI-15). Serverdienst api.php mit zweiter Rennliste.','Nutzer 26.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
