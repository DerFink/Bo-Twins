import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.25.xlsx").replace("0.18.xlsx","0.24.xlsx").replace("Version 0.19","Version 0.25").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.1'
ws=wb['Deckblatt']; ws['B4']='0.25'; ws['B5']='09.10.2026'; ws['B6']=ws['B6'].value.replace('Build 2026-10-08.14',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
i=idx(rows,'STE-21')+1
rows.insert(i,['STE-22','Steuerung und Physik','Start: Bodencrew hält den Ballon','Am Start hält die Bodencrew den Ballon am Boden fest, bis er eindeutig leichter als Luft ist: Erst bei rund 30 kg Auftriebsüberschuss (zwei Säcke) lässt sie los. Bis dahin steht der Ballon still, wird nicht vom Wind über den Boden gezogen, und die Fahrzeit läuft nicht. Ein Hinweis nennt, wie viele kg noch abzuwerfen sind; beim Loslassen kommt eine Meldung. Auch die Vergleichsfahrten und die Bestweite starten mit diesem Überschuss.','Mit weniger als 30 kg Abwurf bleibt der Ballon am Startplatz stehen (Weg 0 m); nach dem zweiten Sack hebt er ab und steigt zügig; der Hinweis zählt die fehlenden kg herunter.','Muss','Umgesetzt',N+': vorher wurde der Ballon bei zu wenig Abwurf über den Boden gezogen',None])
r=rows[idx(rows,'KRT-04')]; r[3]=(r[3] or '')+' Zoomt der Spieler selbst, bleibt sein Ausschnitt erhalten: Die Kamera stellt den Zoom dann nicht mehr von sich aus ein und hält nur den Ballon in der Mitte, bis „Zentrieren“ angetippt wird. Das gilt auch während der Kartenfahrt vom Ziel zum Start.'; r[7]=(r[7] or '')+'; '+N+': Fehler behoben (Karte sprang nach eigenem Zoomen zurück)'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-08.14',BLD)
ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
for c,x in enumerate(('Start: Auftriebsüberschuss, bei dem die Bodencrew loslässt',30,'kg','Vorgabe Nutzer 09.10.2026: rund 30 kg, zwei Säcke (STE-22)'),1): ws.cell(r0,c).value=x; ws.cell(r0,c)._style=copy.copy(ds[c-1])
ws.cell(r0,2).number_format='General'
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-22')]; r[2]=r[2]+', STE-22'; r[5]=r[5]+'; am Start erst einen Sack, dann den zweiten abwerfen'; r[6]=r[6]+'; mit einem Sack bleibt der Ballon stehen, nach dem zweiten lässt die Crew los'
r[11]=(r[11] or '')+f' Build {BLD}: mit 10 und 25 kg Abwurf bleibt der Ballon stehen (Weg 0 m), bei 30 kg hebt er ab; Hinweis und Meldung erscheinen.'
r=rows[idx(rows,'T-30')]; r[5]=r[5]+'; im Zielfahren gleich nach dem Start und in der Fahrt mit zwei Fingern hinein- und herauszoomen, danach „Zentrieren“ antippen'; r[6]=r[6]+'; der eigene Zoom bleibt erhalten, die Karte springt nicht zurück; nach „Zentrieren“ führt die Kamera wieder selbst'
r[11]=(r[11] or '')+f' Build {BLD}: Zoomlogik am Rechner nur mit einer Karten-Attrappe geprüft, nicht mit der echten Karte und nicht mit Fingergesten.'
put(ws,5,rows,14,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.25','09.10.2026',f'Build {BLD}: Am Start hält die Bodencrew den Ballon, bis er 30 kg leichter als Luft ist (STE-22, neuer Parameter, T-22). Karte: eigener Zoom des Spielers bleibt erhalten, kein Zurückspringen (KRT-04, T-30).','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
