import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.49.xlsx").replace("0.18.xlsx","0.48.xlsx").replace("Version 0.19","Version 0.49").replace("Nutzer 08.10.2026","Nutzer 12.10.2026"))
BLD='2026-10-12.2'
ws=wb['Deckblatt']; ws['B4']='0.49'; ws['B6']=ws['B6'].value.replace('Build 2026-10-12.1',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def ins(pre,row):
    i=max(k for k,r in enumerate(rows) if str(r[0]).startswith(pre+'-'))+1; rows.insert(i,row)
ins('STE',['STE-30','Steuerung und Physik','Autopilot','Der Knopf „AP“ unter dem Tau hält die Höhe, die der Ballon beim Einschalten hat, auf ±5 m (Wunsch Nutzer). Darunter wählt ein Schalter Höhe NN oder Höhe über Grund; über Grund erst ab 100 m. Gesteuert wird alle 2 s Fahrzeit mit kleinen Portionen Sand oder Gas: Der Autopilot bestimmt den Auftriebsüberschuss, der die Steiggeschwindigkeit in 60 s auf den Sollwert bringt (Sollwert: Abweichung durch 40 s, höchstens 1 m/s; innerhalb von 2 m keine Höhenkorrektur, nur Dämpfung). Er schaltet sich ab bei eigenem Sand oder Ventil, wenn der Sand unter die Landereserve fällt (6 kg je 100 m³ Hülle, beim Standardballon 60 kg) und wenn der Füllgrad unter 40 % liegt; eine Meldung nennt den Grund. Die Sollhöhe steht als violette Linie „AP“ im Vertikalprofil; bei eingefahrenem Bedienfeld zeigt die Kurzfassung „AP“. Erlaubt in allen Spielarten; eine Liste im Code (AP_HIDE) kann ihn für einzelne Aufgaben ausblenden. Im Tutorial ist er ausgeblendet.','Im Wetter des Rennens über 10 h in 1000 m: größte Abweichung 4,2 m, 20,5 kg Sand und 6,5 % Gas, gegenüber 18,6 kg und 6,3 % bei einem Band von ±30 m; in der App eingeschaltet im Steigen mit 2 m/s auf ±3 m abgefangen.','Kann','Umgesetzt','Wunsch Nutzer 12.10.2026; Grenzen von Claude vorgeschlagen, vom Nutzer nicht beanstandet',None,'Feature'])
r=rows[idx(rows,'NFA-08')]; r[3]=(r[3] or '')+' Seit Build '+BLD+' erklärt das Tutorial die Griffe: Anzeigen (Seite 4) und Bedienfeld (Seite 2).'; r[7]=(r[7] or '')+'; '+N+': Findings Nutzer 12.10.2026'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-12.1',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.49','12.10.2026',f'Build {BLD}: Autopilot mit ±5 m, NN oder über Grund (STE-30, neu). Tutorial erklärt die Griffe von Anzeigen und Bedienfeld (NFA-08).','Nutzer 12.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
