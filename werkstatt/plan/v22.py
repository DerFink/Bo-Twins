import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.22.xlsx").replace("0.18.xlsx","0.21.xlsx").replace("Version 0.19","Version 0.22"))
BLD='2026-10-08.11'
ws=wb['Deckblatt']; ws['B4']='0.22'; ws['B6']=ws['B6'].value.replace('Build 2026-10-08.10',f'Build {BLD}')
ws['B8']='Web-App für Android und iOS (iPhone, iPad) im Browser, installierbar auf dem Startbildschirm; bisher im Ordner /spiel/ auf GitHub Pages, jetzt auch als Paket für einen eigenen Server (SRV-01)'
ws=wb['Ziele']; rows=rows_of(ws,5,4)
i=idx(rows,'Z-12',1)+1
rows.insert(i,['Ziel','Z-13','Die richtige Physik bleibt unbedingt erhalten; sie ist der Kern des Programms. Dazu gehört, dass ein praller Ballon beim Aufstieg über seine Gleichgewichtshöhe hinausschießt, Gas abbläst und danach wieder sinkt, je nach Luftschichtung und Temperaturgradient. Es gibt keine vereinfachte Fahrweise.',N])
put(ws,5,rows,4,4)

ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
def a(key,**kw):
    m={'C':2,'D':3,'E':4,'F':5,'G':6,'H':7}; r=rows[idx(rows,key)]
    for k,v in kw.items(): r[m[k]]=v
r=rows[idx(rows,'STE-05')]; r[3]=r[3]+' Der Ballon kommt mit Restfahrt an seiner Gleichgewichtshöhe an und schießt wegen seiner Masse darüber hinaus (nach 30 kg Abwurf rund 30 m); dabei bläst er weiteres Gas ab, ist danach zu schwer und sinkt. Ob er sich fängt oder bis zum Boden durchfällt, hängt von der Luftschichtung ab; abfangen lässt er sich mit Ballast. Dieses Verhalten ist gewollt und bleibt (Z-13).'; r[7]=(r[7] or '')+'; '+N+': richtige Physik beibehalten'
r=rows[idx(rows,'STE-08')]; r[3]=r[3]+' Links neben dem Vario zeigt ein Balken die Vertikalbeschleunigung: Bei 0 m/s² ist er nicht zu sehen; bei positiver Beschleunigung wächst er grün nach oben, bei negativer rot nach unten. Die Länge folgt einer Wurzelskala, Vollausschlag bei 0,25 m/s², damit auch kleine Werte sichtbar sind.'; r[4]='Alle genannten Werte sind sichtbar und aktualisieren sich; nach einem Sandabwurf wächst der Balken grün nach oben, nach einem Ventilzug rot nach unten, im gleichmäßigen Steigen oder Sinken und am Boden ist er nicht zu sehen.'; r[7]=(r[7] or '')+'; '+N+': Balken der Vertikalbeschleunigung'
a('INS-01',D='Das Spiel besteht aus einem Ordner mit eigenem Manifest, eigenen Symbolen und eigenem Service Worker; alle Pfade sind relativ zu diesem Ordner. Es läuft deshalb unter jeder Adresse (bisher /spiel/ auf GitHub Pages, künftig auf einem eigenen Server) und lässt sich getrennt von der Tracker-App installieren.')
a('INS-03',D='Eine neue Version entsteht durch Hochladen der geänderten Dateien auf den Server. Die App fragt beim Start mit Netz nach der neuesten Fassung, lädt sie und meldet „Neue Version geladen“; ohne Netz startet die zuletzt geladene Fassung. Jede Fassung trägt eine Build-Nummer, sichtbar im Startfenster und in den Einstellungen; ein Knopf erzwingt das Laden der neuesten Version. Bei jedem Build wird die Versionsnummer im Service Worker erhöht.',E='Nach dem Hochladen zeigt die App beim nächsten Start mit Netz die neue Build-Nummer; ohne Netz startet die alte Fassung weiter.')
a('SRV-01',G='Teilweise',H=N+': Paket für den Server liegt vor (Gasballon_WebApp), mit Anleitung; noch nicht auf dem Server des Nutzers geprüft')
i=idx(rows,'INS-03')+1
rows.insert(i,['INS-04','Installation','Hinweis zur Installation auf Android und iOS','Im Startfenster erscheint ein Hinweis „Als App installieren“ mit den Knöpfen „Installieren“ und „Später“. Auf Android öffnet „Installieren“ den Installationsdialog des Browsers; bietet der Browser keinen an, nennt der Hinweis den Weg über das Browsermenü. Auf iPhone und iPad nennt er die Schritte in Safari (Teilen, Zum Home-Bildschirm, Hinzufügen). „Später“ blendet den Hinweis für 7 Tage aus; über einen Knopf in den Einstellungen lässt er sich jederzeit wieder aufrufen. In der installierten App erscheint er nicht. Der Hinweis ist in allen sieben Sprachen vorhanden. Die installierte App startet auch ohne Netz.','Auf einem Android-Handy und einem iPhone erscheint der Hinweis im Browser; die Installation gelingt nach den genannten Schritten; die installierte App zeigt keinen Hinweis und startet im Flugmodus.','Muss','Umgesetzt',N+'; Vorbild cloud.aerarium.de/doppelkopf',None])
rows.append(['SRV-06','Server','Aktualisierungen und Events','Nach der Veröffentlichung lassen sich Aktualisierungen einspielen, zum Beispiel ein Event (eine besondere Aufgabe für einen Zeitraum). Der Ablauf: Änderung mit Claude bauen und prüfen, neue Build-Nummer, geänderte Dateien auf den Server laden, auf einem Gerät die Build-Nummer prüfen. Offen ist, ob Events als neuer Build kommen oder als eigene Datei auf dem Server, die sich ohne neuen Build ändern lässt (OP-42).','Eine hochgeladene Änderung ist bei allen Nutzern beim nächsten Start mit Netz aktiv; ein Event ist zum angekündigten Zeitpunkt spielbar.','Soll','Offen',N,None])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA)
ws['A2']=ws['A2'].value.replace('2026-10-08.10',BLD)

ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
for c,x in enumerate(('Vario: Vollausschlag des Beschleunigungsbalkens',0.25,'m/s²','Vorschlag Claude; Wurzelskala. Ein Sack von 15 kg ergibt rund 0,08 m/s², 30 kg rund 0,16 (STE-08)'),1): ws.cell(r0,c).value=x; ws.cell(r0,c)._style=copy.copy(ds[c-1])
ws.cell(r0,2).number_format='General'

ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-16')]; r[2]='INS-01, INS-02, INS-03, INS-04, SRV-01, SRV-06, R-10, R-18'; r[4]='Android-Handy und iPhone mit Netz und Flugmodus; App auf dem Server'
r[5]='App im Browser öffnen; Hinweis zur Installation lesen, „Später“ und über die Einstellungen wieder aufrufen; auf Android und auf iPhone installieren; installierte App im Flugmodus starten; neue Version hochladen und zweimal starten'
r[6]='Der Hinweis erscheint auf beiden Systemen und führt zur Installation; die installierte App zeigt keinen Hinweis und startet im Flugmodus; nach dem Hochladen zeigt sie beim nächsten Start mit Netz die neue Build-Nummer'
r[11]=f'Build {BLD}: Am Rechner mit nachgestelltem Android und iPhone geprüft: Hinweis erscheint, „Installieren“ öffnet den Dialog des Browsers (Android) oder zeigt die Schritte (iPhone), „Später“ wirkt, installierte App ohne Hinweis. Echte Installation und Start ohne Netz nur am Gerät prüfbar.'
r=rows[idx(rows,'T-08')]; r[5]=r[5]+'; nach Sandabwurf und nach Ventilzug den Balken am Vario ansehen'; r[6]=r[6]+'; der Balken am Vario wächst bei Beschleunigung nach oben grün, nach unten rot und verschwindet bei gleichmäßiger Fahrt'
r[11]=r[11]+f' Build {BLD}: Balken nach 30 kg Abwurf grün (0,16 m/s²), nach Gasverlust rot (−0,18 m/s²), am Boden nicht sichtbar.'
r=rows[idx(rows,'T-06')]; r[6]=r[6].replace('und der Ballon steigt nicht unbegrenzt weiter;','und der Ballon steigt nicht unbegrenzt weiter; er schießt über die Gleichgewichtshöhe hinaus und sinkt danach, bis er abgefangen wird oder die Luftschichtung ihn trägt;')
put(ws,5,rows,14,4)

ws=wb['Risiken']; rows=rows_of(ws,6,15)
rows.append(['R-18','Technik','Die App verhält sich auf iPhone und iPad anders als auf Android','Bisher nur für Android gebaut und nur am Rechner geprüft; Safari unterscheidet sich bei Ton, Installation, Speicher und Standort','Töne fehlen, Installation oder Offline-Start misslingen, gespeicherte Daten gehen nach längerer Pause verloren',3,3,None,None,'Test auf einem iPhone vor der Bekanntgabe der Adresse (T-16, T-26)','Hinweis für iOS anpassen; Fehler einzeln beheben','T-16','Nutzer, Claude','Offen','INS-04, SRV-01'])
def fixR(v,r): v[7]=f'=F{r}*G{r}'; v[8]=f'=IF(H{r}>=15,"hoch",IF(H{r}>=8,"mittel","niedrig"))'
put(ws,6,rows,15,5,fixR)

ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
def o(key,**kw):
    m={'C':2,'D':3,'F':5,'G':6,'H':7,'I':8,'J':9}; r=rows[idx(rows,key)]
    for k,v in kw.items(): r[m[k]]=v
o('OP-34',H='Der Nutzer hat einen Server zum Ablegen der Dateien (Nutzer 08.10.2026); das Paket für den Server liegt vor. Offen: Adresse, und ob der Server Programme ausführen kann (PHP mit Datenbank oder Ähnliches) für Bestenliste und Statistik',J='In Klärung')
rows+=[['OP-42','Server','Wie sollen Events eingespielt werden: als neuer Build (jede Änderung über Claude, dafür ohne zusätzliche Technik) oder als eigene Datei auf dem Server, die Aufgabe, Zeitraum, Startplatz und Wetter eines Events beschreibt und sich ohne neuen Build ändern lässt? Was ist ein Event genau (gleiche Aufgabe und gleiches Wetter für alle, eigene Bestenliste, Zeitraum)?','Neuer Build je Event / Event-Datei auf dem Server','Nutzer','12','B',None,None,'Offen'],
       ['OP-43','Physik','Nach dem Aufstieg schießt der pralle Ballon über die Gleichgewichtshöhe hinaus, bläst Gas ab und sinkt wieder. Soll das vereinfacht werden?','Richtige Physik / vereinfachte Fahrweise','Nutzer','2','A','Entschieden (Nutzer 08.10.2026): Die richtige Physik bleibt unbedingt erhalten (Z-13, STE-05). Sinken nach dem Steigen gehört dazu und hängt von Luftschichtung und Temperaturgradient ab. Die simulierte Luft bleibt unverändert (6 bis 9,5 K/km unter der Inversion)','08.10.2026','Geschlossen']]
put(ws,5,rows,10,4)

ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.22','08.10.2026',f'Build {BLD}: Balken der Vertikalbeschleunigung am Vario (STE-08). Hinweis zur Installation auf Android und iOS (INS-04), Symbol und Angaben für iPhone und iPad. Paket für den eigenen Server mit Anleitung (SRV-01, INS-01, INS-03). Ablauf für Aktualisierungen und Events (SRV-06, OP-42). Entscheidung: richtige Physik bleibt, kein vereinfachtes Einpendeln (Z-13, STE-05, OP-43). Neues Risiko R-18 (iOS ungetestet).','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert')
