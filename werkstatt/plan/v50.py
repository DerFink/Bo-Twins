import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.50.xlsx").replace("0.18.xlsx","0.49.xlsx").replace("Version 0.19","Version 0.50").replace("Nutzer 08.10.2026","Nutzer 13.10.2026"))
BLD='2026-10-13.1'
ws=wb['Deckblatt']; ws['B4']='0.50'; ws['B6']=ws['B6'].value.replace('Build 2026-10-12.2',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'STE-22')]; r[3]=(r[3] or '')+' Seit Build '+BLD+': Die Untergrenze hängt vom Volumen ab, je angefangene 500 m³ ein Sack (400 m³: 15 kg, 1000 m³: 30 kg, 1500 m³: 45 kg); Hinweis und Tutorial nennen den tatsächlichen Wert. Gerechnet im simulierten Wetter: 400 m³ steigt mit einem Sack auf rund 380 m, 1000 m³ mit zwei auf rund 310 m, 1500 m³ mit drei auf rund 315 m.'; r[7]=(r[7] or '')+'; '+N+': Finding Nutzer 13.10.2026 (400 m³ ist eine gängige Größe)'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-12.2',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.50','13.10.2026',f'Build {BLD}: Startschwelle der Bodencrew abhängig vom Volumen (STE-22). Ab jetzt legt Claude freigegebene Builds selbst auf GitHub (spiel/) ab; Werkzeuge liegen in werkstatt/.','Nutzer 13.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
