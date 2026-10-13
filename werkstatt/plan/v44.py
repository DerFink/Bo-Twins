import openpyxl, copy, re
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.44.xlsx").replace("0.18.xlsx","0.43.xlsx").replace("Version 0.19","Version 0.44").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
ws=wb['Deckblatt']; ws['B4']='0.44'
ws=wb['Tests']; rows=rows_of(ws,5,15); done=[]
for r in rows:
    m=re.match(r'T-(\d+)$',str(r[0]))
    if m and 1<=int(m.group(1))<=14 and r[0] not in ('T-06','T-13') and r[10]!='Bestanden':
        r[10]='Bestanden'; r[11]=(r[11] or '')+' Nutzer 11.10.2026: am Handy bestanden.'; r[12]='11.10.2026'; r[13]='Nutzer'; done.append(r[0])
r=rows[idx(rows,'T-06')]; r[11]=(r[11] or '')+' Nutzer 11.10.2026: Das zu prüfen ist eine eigene Messkampagne und dauert; bleibt offen.'
r=rows[idx(rows,'T-13')]; r[11]=(r[11] or '')+' Nutzer 11.10.2026: Es gibt kein Fahrzeug; „wäre aber lustig, wenn es gezeigt würde“. Bleibt zurückgestellt.'
put(ws,5,rows,15,4); tests=rows; print('bestanden gesetzt:',done)
ws=wb['Anforderungen']; rows=rows_of(ws,5,10); up=[]
for r in rows:
    if r[6]!='Umgesetzt': continue
    ref=[t for t in tests if re.search(r'(?<![A-Z0-9-])'+re.escape(str(r[0]))+r'(?!\d)',str(t[2]))]
    if ref and all(t[10]=='Bestanden' for t in ref): r[6]='Erledigt'; up.append(r[0])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); print('erledigt:',up)
ab=wb['Abnahme_Basis']
for r in range(ab.max_row,4,-1):
    if ab.cell(r,1).value in done: ab.delete_rows(r)
    elif ab.cell(r,1).value=='T-06': ab.cell(r,7).value='Eigene Messkampagne, dauert (Nutzer 11.10.2026)'
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.44','11.10.2026','Abnahme: '+', '.join(done)+' am Handy bestanden (Nutzer). T-06 bleibt offen (Messkampagne), T-13 zurückgestellt. Anforderungen, deren Tests alle bestanden sind, stehen auf „Erledigt“: '+', '.join(up)+'. Kein neuer Build.','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT)
