import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.32.xlsx").replace("0.18.xlsx","0.31.xlsx").replace("Version 0.19","Version 0.32").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.8'
ws=wb['Deckblatt']; ws['B4']='0.32'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.7',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
def app(key,txt,src=None):
    r=rows[idx(rows,key)]; r[3]=(r[3] or '')+' '+txt; r[7]=(r[7] or '')+'; '+N+(': '+src if src else '')
r=rows[idx(rows,'PUN-01')]; r[2]='Punkte nach Abstand, Landezone 100 m Durchmesser'; r[3]=r[3].replace('100 m Radius','100 m Durchmesser (Radius 50 m)'); r[7]=(r[7] or '')+'; '+N+': Durchmesser, außerhalb linear fallend (OP-41)'
r=rows[idx(rows,'PUN-02')]; r[2]='Abzug und Disqualifikation bei hartem Aufprall'; r[3]='Bis 1,0 m/s Sinken beim Aufsetzen gibt es im Zielfahren keinen Abzug, darüber fällt die Punktzahl linear. Wer härter als 2,5 m/s aufsetzt, ist disqualifiziert: 0 Punkte in Zielfahren, Wandziel, MDT und MID, in der Meisterschaft keine Wertung. Die Auswertung nennt den Grund.'; r[4]='Landung mit 2,4 m/s wird gewertet, mit 2,6 m/s nicht.'; r[7]=(r[7] or '')+'; '+N+': Entscheidung zu OP-05'
app('MEI-01','Der Satz lautet jetzt: „Ein Team schied leider aus.“','Text gekürzt')
app('MEI-03','Um die Grenzen der gesperrten Länder liegt ein roter Streifen von 11 km Breite: die 10 km Abstand der Regel und 1 km Reserve. Wer den Streifen nicht berührt, ist sicher.','Sperrzone sichtbar größer')
app('PUN-07','Unter dem Namensfeld steht zusätzlich: „Der Name wird ausschließlich für die Bestenliste verwendet und dort gespeichert.“','Entscheidung zu OP-35')
i=idx(rows,'MEI-14')+1
rows.insert(i,['MEI-15','Meisterschaft','Platz jetzt','Während der Fahrt zeigt die Anzeige den Platz, den man hätte, wenn man jetzt landen würde (zum Beispiel 14 / 24), gerechnet aus der aktuellen Entfernung zum Startplatz gegen die Ergebnisse der echten Teams.','Der Platz steigt mit der Entfernung und stimmt nach der Landung mit der Auswertung überein.','B','Umgesetzt',N,None])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.7',BLD)
ws=wb['Parameter']
ws.cell(51,1).value='Zielfahren: Durchmesser der Landezone (volle Punktzahl)'; ws.cell(51,4).value='Entscheidung Nutzer 09.10.2026: Durchmesser 100 m, also Radius 50 m; außerhalb linear fallend bis zum Punkteradius (mindestens 300 m) (PUN-01, OP-41)'
ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
for k,v in enumerate([('Landung: härter als dieser Wert gilt als disqualifiziert',2.5,'m/s','Entscheidung Nutzer 09.10.2026 (PUN-02, OP-05)'),('Meisterschaft: Reserve, um die der rote Streifen breiter gezeichnet ist als die Regel',1,'km','Vorgabe Nutzer 09.10.2026 (MEI-03)')]):
    for c,x in enumerate(v,1): ws.cell(r0+k,c).value=x; ws.cell(r0+k,c)._style=copy.copy(ds[c-1])
    ws.cell(r0+k,2).number_format='General'
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-37')]; r[2]=r[2]+', MEI-15'; r[5]=r[5]+'; auf die Zeile „Platz jetzt“ achten; die Breite des roten Streifens an einer Grenze mit dem Kartenmaßstab vergleichen'
r[6]=r[6]+'; der Platz verbessert sich mit der Entfernung; der Streifen ist rund 11 km breit und in jeder Zoomstufe gleich breit in km'
r[11]=(r[11] or '')+f' Build {BLD}: „Platz jetzt“ erscheint nach dem Loslassen (24 / 24 bei 197 km); Landung mit 3,7 m/s wird als disqualifiziert geführt; Wertung Zielfahren: 40 m 100 Punkte, 60 m 98, mit 2,6 m/s 0. Der rote Streifen ließ sich am Rechner nicht auf einer echten Karte ansehen; Breite und Aussehen am Handy prüfen.'
put(ws,5,rows,14,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
def close(k,txt,st='Geschlossen'):
    r=rows[idx(rows,k)]; r[7]=((r[7]+' ') if r[7] else '')+txt; r[8]='09.10.2026'; r[9]=st
close('OP-48','Wetterpaket über die App erzeugt (Nutzer, 09.10.2026); daraus sind zwei Routen gerechnet (MEI-11).')
close('OP-41','Durchmesser 100 m, außerhalb linear fallend. Umgesetzt in Build '+BLD+'.')
close('OP-05','Härter als 2,5 m/s gilt als disqualifiziert; darunter bleibt der lineare Abzug im Zielfahren. Umgesetzt in Build '+BLD+' für alle gewerteten Aufgaben und die Meisterschaft.')
close('OP-31','So lassen.')
close('OP-28','Ist so in Ordnung: ein Wert für alle Flächen.')
close('OP-12','Overpass bleibt (öffentlicher Abfragedienst für OpenStreetMap); die eigene Ortsliste entfällt.')
close('OP-35','Veröffentlicht werden nur die selbst gewählten Namen in der Bestenliste; ein Hinweis steht unter dem Namensfeld. Solange es keinen gemeinsamen Server gibt, bleiben die Namen auf dem Gerät. Keine Rechtsberatung; bei einer gemeinsamen Liste auf dem Server noch einmal ansehen.')
close('OP-46','Nutzer fragt erfahrene Gasballonpiloten.','In Klärung')
close('OP-34','Synology NAS. Für die App selbst reicht das Ausliefern der Dateien über HTTPS. Für eine gemeinsame Bestenliste müsste auf dem NAS ein kleines Programm laufen (Web Station mit PHP und Datenbank).','In Klärung')
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.32','09.10.2026',f'Build {BLD}: Meisterschaft zeigt den „Platz jetzt“ (MEI-15) und einen roten Streifen von 11 km um die gesperrten Länder (MEI-03); Text gekürzt (MEI-01). Landezone im Zielfahren 100 m Durchmesser (PUN-01); härter als 2,5 m/s gelandet gilt als disqualifiziert (PUN-02); Hinweis zur Verwendung des Namens (PUN-07). Geschlossen: OP-05, OP-12, OP-28, OP-31, OP-35, OP-41, OP-48; Antworten zu OP-34 und OP-46 eingetragen.','Nutzer 09.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
