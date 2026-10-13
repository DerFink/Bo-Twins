import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.46.xlsx").replace("0.18.xlsx","0.45.xlsx").replace("Version 0.19","Version 0.46").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
BLD='2026-10-11.4'
ws=wb['Deckblatt']; ws['B4']='0.46'; ws['B6']=ws['B6'].value.replace('Build 2026-10-11.3',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'STE-29')]; r[3]=(r[3] or '')+' Seit Build '+BLD+': Ist die Hülle leer (Füllgrad zeigt 0 %), bleibt das Ventil stumm und in der Seitenansicht tritt kein Gas mehr aus.'; r[7]=(r[7] or '')+'; '+N+': Fund Nutzer, bei 0 % Gas kam noch das Ventilgeräusch'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-11.3',BLD)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.46','11.10.2026',f'Build {BLD}: Bei leerer Hülle kein Ventilgeräusch und kein austretendes Gas mehr (STE-29).','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
