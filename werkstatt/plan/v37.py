import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.37.xlsx").replace("0.18.xlsx","0.36.xlsx").replace("Version 0.19","Version 0.37").replace("Nutzer 08.10.2026","Nutzer 10.10.2026"))
BLD='2026-10-10.4'
ws=wb['Deckblatt']; ws['B4']='0.37'; ws['B6']=ws['B6'].value.replace('Build 2026-10-10.3',f'Build {BLD}')
ws=wb['Anforderungen']; n0=ws.max_row; rows=rows_of(ws,5,10)
def app(key,txt,src=None):
    r=rows[idx(rows,key)]; r[3]=(r[3] or '')+' '+txt; r[7]=(r[7] or '')+'; '+N+(': '+src if src else '')
r=rows[idx(rows,'STE-27')]; r[2]='Temp als Skew-T-Diagramm'; r[3]='Unter dem Vertikalprofil sitzt der Knopf „Temp“. Er blendet ein Skew-T-Diagramm ein, wie es die Wetterdienste zeigen; ein Tipp auf das Diagramm blendet es aus. Druck logarithmisch nach oben, Isothermen schräg, Trockenadiabaten orange, Feuchtadiabaten grün, Mischungsverhältnis gepunktet; Temperatur rot, Taupunkt blau (nur mit echtem Wetter, aus der relativen Feuchte je Druckfläche); Inversionen als hellblaues Band mit Höhenangabe, Wolkendecke als weißer Strich mit Schrift, Windfahnen wie im Vertikalprofil, Höhe in km, Marke für Ballon und Gastemperatur. Das Vertikalprofil selbst bleibt fest und reagiert nicht mehr auf Antippen.'; r[4]='Der Knopf blendet das Diagramm zuverlässig ein; es zeigt Temperatur und Taupunkt wie der Temp des Wetterdienstes.'; r[7]=(r[7] or '')+'; '+N+': Fund „nicht reproduzierbar aufrufbar“, Vorlage Temp Essen; Stüve-Fassung ersetzt'
app('STE-28','Solange die Seitenansicht sichtbar ist, folgt die Karte dem Ballon; der Knopf „Ballon folgen“ erscheint dann nicht.','Fund: Knopf lag über der Seitenansicht')
app('SPI-17','Der Sand im Bild fällt nur bis zum Boden. Auf niedrigen Bildschirmen zeigen die Menüpunkte nur ihre Titel. Unter „Buy me a Beer“ steht „Als App installieren“, solange der Hinweis zur Installation zurückgestellt ist.','Funde vom Handy')
app('SPI-18','Jetzt 16 Schritte: neu sind das Ein- und Ausblenden des Temp und der Hinweis auf die Einstellungen. „Weiter“ ist immer zu sehen und ausgegraut, solange eine Handlung erwartet wird; der Knopf „Beenden“ entfällt (zurück geht es über das Hauptmenü). Texte ergänzt: Flattern des Parachute-Ventils, Schlepptau erst unter 50 m sinnvoll.','Funde vom Handy')
app('TON-02','Nach dem Fund „auf dem Handy nicht gehört“ eine Oktave höher und rund dreimal so laut: Handy-Lautsprecher geben die tiefen Töne der ersten Fassung kaum wieder.','Fund vom Handy')
app('MET-01','Zusätzlich wird die relative Feuchte je Druckfläche und in 2 m geladen (nur in der erweiterten Abfrage); daraus rechnet die App den Taupunkt für den Temp.','Taupunkt')
app('INS-04','Der Hinweis erscheint nach „Später“ sieben Tage nicht; in dieser Zeit steht auf der Startseite der Menüpunkt „Als App installieren“.','Fund: Frage nicht mehr gesehen')
i=max(k for k,r in enumerate(rows) if str(r[0]).startswith('NFA-'))+1
rows.insert(i,['NFA-07','Nichtfunktional','Rollhinweis','Wo ein Fenster länger ist als der Bildschirm (Startseite, Einstellungen, Startfenster, Vorbereitung, Karten und Daten), liegt am unteren Rand ein oranger Pfeil, bis das Ende erreicht ist.','Der Pfeil erscheint nur, wenn es nach unten weitergeht, und verschwindet am Ende.','Soll','Umgesetzt',N+': Fund „untere Hälfte nicht sichtbar, man muss scrollen“',None,'Basis'])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-10.3',BLD)
ws=wb['Parameter']
for i in range(1,ws.max_row+1):
    v=str(ws.cell(i,1).value or '')
    if 'GPS' in v and ws.cell(i,2).value==10: ws.cell(i,2).value=15; ws.cell(i,4).value=(ws.cell(i,4).value or '')+'; Nutzer 10.10.2026: 15 s'
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-40')]; r[5]=r[5].replace('links aufs Profil tippen, Diagramm ansehen, antippen','unter dem Profil auf „Temp“ tippen, Diagramm ansehen, antippen'); r[6]=r[6].replace('Stüve-Diagramm öffnet und schließt per Tipp','der Temp (Skew-T) öffnet über den Knopf und schließt per Tipp; mit echtem Wetter ist der Taupunkt zu sehen'); r[11]=(r[11] or '')+f' Build {BLD}: Skew-T mit nachgestellter Feuchte geprüft (Taupunkt 7,7 °C am Boden, Inversionsband, Wolkenstrich); Karte folgt bei sichtbarer Seitenansicht. Offen am Handy: Abruf der Feuchte am echten Dienst.'
r=rows[idx(rows,'T-41')]; r[2]=r[2]+', NFA-07'; r[6]=r[6]+'; „Weiter“ ist immer zu sehen; der Rollpfeil zeigt, wo es weitergeht'; r[11]=(r[11] or '')+f' Build {BLD}: 16 Schritte durchlaufen, Rollpfeil und Installationslink geprüft (Bildschirm 360 × 640). Melodie weiterhin nicht anhörbar.'
put(ws,5,rows,15,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'OP-61')]; r[7]='Nutzer 10.10.2026: Skew-T wie der Temp des Wetterdienstes, mit Taupunkt. Umgesetzt in Build '+BLD+'.'; r[8]='10.10.2026'; r[9]='Geschlossen'
r=rows[idx(rows,'OP-60')]; r[7]='Fund 10.10.2026: auf dem Handy nicht zu hören. Build '+BLD+': eine Oktave höher und lauter. Erneut anhören.'; r[9]='In Klärung'
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.37','10.10.2026',f'Build {BLD}: Funde vom Handy. Temp als Skew-T mit Taupunkt über einen Knopf unter dem Profil (STE-27, MET-01, OP-61 geschlossen). Melodie höher und lauter (TON-02, OP-60). Tutorial mit 16 Schritten, „Weiter“ immer sichtbar (SPI-18). Rollhinweis (NFA-07). Installationslink auf der Startseite (INS-04). Karte folgt bei sichtbarer Seitenansicht (STE-28). GPS-Mittelung 15 s.','Nutzer 10.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
