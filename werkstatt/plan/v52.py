import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.52.xlsx").replace("0.18.xlsx","0.51.xlsx").replace("Version 0.19","Version 0.52").replace("Nutzer 08.10.2026","Nutzer 26.10.2026"))
BLD='2026-10-26.4'
ws=wb['Deckblatt']; ws['B4']='0.52'; ws['B6']=ws['B6'].value.replace('Build 2026-10-26.1',f'Build {BLD}').replace('Build 2026-10-13.1',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def ins(pre,row):
    i=max(k for k,r in enumerate(rows) if str(r[0]).startswith(pre+'-'))+1; rows.insert(i,row)
r=rows[idx(rows,'MEI-14')]; r[3]=(r[3] or '')+' 26.10.2026: Das Open-Meteo-Archiv liefert für 2018 keine Werte auf Druckflächen (alle null). Seit Build 2026-10-26.2 rechnet die App fehlende Höhen der Druckflächen mit der barometrischen Höhenformel und nennt bei Fehlschlag den Grund; seit Build 2026-10-26.3 ist 2018 als „in Vorbereitung“ gesperrt, bis ein Wetterpaket aus ERA5 (Copernicus) vorliegt.'; r[6]='In Arbeit'
ins('STE',['STE-31','Steuerung und Physik','Einheitliches Ballonbild','Ein Bild für App-Symbol, Logo, Kartensymbol, Seitenansicht und Startseite (Entwurf B2, Wahl Nutzer 26.10.2026): silberne Hülle mit orangem Rautennetz (fünf Reihen), oranges Ventil, Füllansatz, Lastring, Weidenkorb mit zwei Sandsäcken. Die Form folgt dem Füllgrad wie bisher (STE-29). Eine Geometrie (balGeo) für SVG und Leinwand; App-Symbole 192, 512, maskierbar 512 und Apple 180 daraus erzeugt.','Alle fünf Stellen zeigen dieselbe Zeichnung; auf der Karte in Originalgröße auf heller und dunkler Karte erkennbar.','Soll','Umgesetzt','Wunsch Nutzer 26.10.2026, drei Entwürfe und zwei Netzvarianten vorab gezeigt',None,'Basis'])
r=rows[idx(rows,'INS-06')]; r[3]=(r[3] or '')+' Seit Build '+BLD+': Nach einem Update erscheint nur ein Satz („Die App wurde aktualisiert, Neues steht unter ‚Was ist neu‘.“), die Liste öffnet sich nicht mehr von selbst. In der Liste bleibt der Knopf „Schließen“ unten immer sichtbar.'; r[7]=(r[7] or '')+'; Nutzer 26.10.2026'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA)
import re
ws['A2']=re.sub(r'2026-\d\d-\d\d\.\d',BLD,ws['A2'].value)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.52','26.10.2026',f'Builds 2026-10-26.2 bis {BLD}: Wetter 2018 nicht verfügbar, Szenario gesperrt (MEI-14 In Arbeit). Einheitliches Ballonbild B2 (STE-31, neu). „Was ist neu“ nach Update nur als Satz, Schließen-Knopf immer sichtbar (INS-06).','Nutzer 26.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
