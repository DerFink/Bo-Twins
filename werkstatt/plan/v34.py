import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.34.xlsx").replace("0.18.xlsx","0.33.xlsx").replace("Version 0.19","Version 0.34").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.10'
ws=wb['Deckblatt']; ws['B4']='0.34'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.9',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
for r in rows:
    if r[0] in ('SRV-02','SRV-03','SRV-04'):
        r[3]=(r[3] or '')+' Stand Build '+BLD+': Der Dienst api/api.php (PHP ab 8.0 mit SQLite) liegt dem Paket bei. Die App erkennt ihn beim Start von selbst; ohne Dienst bleibt alles auf dem Gerät. Gemeinsam sind die Bestenliste und die Liste „69. Gordon Bennett“; Besuche werden je Tag gezählt. Gespeichert wird ein Kurzwert der Gerätekennung, keine IP-Adresse.'; r[6]='Teilweise'; r[7]=(r[7] or '')+'; '+N+': Dienst erstellt, auf dem NAS noch nicht in Betrieb'
print([(r[0],r[2],r[6]) for r in rows if str(r[0]).startswith('SRV')])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.9',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-35')]; r[10]='Am Rechner geprüft'; r[11]=(r[11] or '')+f' Build {BLD}: Dienst mit PHP 8.3 auf einem Testrechner geprüft: Lebenszeichen, Ergebnis melden, ungültige Werte abgewiesen (150 Punkte, kaputtes JSON), Begrenzung auf 60 je Stunde, Listen für Challenge und Gordon Bennett, Zurückziehen, Verwaltung nur mit Schlüssel, Löschen, Statistik. Im Browser: App erkennt den Dienst, ein zweites Gerät sieht den Eintrag des ersten, bei Ausfall erscheint die Liste des Geräts mit Hinweis. Offen: Betrieb auf dem Synology-NAS, Schreibrechte, HTTPS, Service Worker.'; r[12]='09.10.2026'; r[13]='Claude (Rechner)'
put(ws,5,rows,14,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
r=rows[idx(rows,'OP-34')]; r[7]=(r[7] or '')+' Build '+BLD+': Serverfassung mit Dienst für die gemeinsame Bestenliste erstellt (Gasballon_Server_Build_'+BLD+'.zip). Wartet auf den Kollegen mit den Admin-Zugängen. Push bei neuen Events ist nicht enthalten (OP-59).'
rows+=[['OP-59','Server','Push-Nachricht und Zeichen am App-Symbol bei neuen Events (Stufe 2): erst bauen, wenn die gemeinsame Bestenliste auf dem NAS läuft. Braucht die Zustimmung der Nutzer, ein Schlüsselpaar und ein Versandskript auf dem Server; auf dem iPhone nur bei installierter App.','Bauen / entfällt','Nutzer','7','C',None,None,'Offen']]
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.34','09.10.2026',f'Build {BLD}: Serverfassung. Dienst api/api.php für gemeinsame Bestenliste, Liste „69. Gordon Bennett“ und Besuchszählung (SRV-02 bis SRV-04, T-35); die App erkennt den Dienst selbst. Startmelodie leiser. Service Worker lässt Aufrufe des Dienstes unberührt (v32). Neu: OP-59 Push.','Nutzer 09.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
