import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.27.xlsx").replace("0.18.xlsx","0.26.xlsx").replace("Version 0.19","Version 0.27").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.3'
ws=wb['Deckblatt']; ws['B4']='0.27'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.2',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
r=rows[idx(rows,'STE-19')]; r[3]=r[3]+' Seit Build 2026-10-09.3 kommen Sonne, Wärmestrahlung und Regen dazu (STE-23, STE-24); die frühere vereinfachte Abstrahlung gegen die Lufttemperatur entfällt.'
r=rows[idx(rows,'SPI-01')]; r[3]=r[3].replace('mittelt die Position etwa 15 Sekunden','mittelt die Position 10 Sekunden; ein großer Countdown zeigt die verbleibenden Sekunden und nennt den Grund des Wartens, während der Suche nach dem Fix steht dort die erreichte Genauigkeit'); r[7]=(r[7] or '')+'; '+N+': 10 statt 15 Sekunden, großer Countdown'
r=rows[idx(rows,'STE-08')]; r[3]=r[3]+' Eine Zeile zeigt den Himmel: Sonne mit Sonnenhöhe in Grad, Nacht, Wolken über dem Ballon oder Regen mit mm/h; bei nasser Hülle kommt eine Zeile mit dem Wasser in kg dazu.'
r=rows[idx(rows,'STE-17')]; r[3]=r[3]+' Wolken erscheinen als heller Streifen neben der Schichtung (je heller, desto dichter), Regen als blaue Punkte darunter.'
i=idx(rows,'STE-22')+1
rows[i:i]=[
 ['STE-23','Steuerung und Physik','Sonne und Nachtabkühlung','Die Hülle ist weiß und aus leichtem, PVC-beschichtetem Gewebe (135 kg auf 484 m², also rund 280 g/m²). Sie nimmt Sonnenlicht auf: direkt auf ihre Querschnittsfläche, dazu Himmelslicht und vom Boden oder von einer Wolkendecke zurückgeworfenes Licht. Der Sonnenstand folgt aus Ort, Datum und Uhrzeit. Die Hülle tauscht außerdem Wärmestrahlung aus: Sie strahlt mit ihrer eigenen Temperatur ab und empfängt die Strahlung des Himmels und des Bodens. Der klare Himmel strahlt schwächer, als die Luft warm ist; nachts kühlt die Hülle deshalb unter die Lufttemperatur, Wolken darüber heben das auf. Das Gas folgt der Hülle. Am Startplatz haben Hülle und Gas schon die Temperatur, die sich in Sonne oder Nacht einstellt; der Trimmballast ist darauf abgewogen. Der Effekt gilt in allen Spielarten.','Gerechnet bei klarem Himmel in 500 m: mittags Gas rund 10 K wärmer als die Luft, nachts rund 6 K kälter; unter Wolken nachts kaum Abkühlung; über den Sonnenuntergang muss Ballast gegeben werden, nach Sonnenaufgang steigt der Ballon von selbst.','Muss','Umgesetzt',N+'; Stoffwerte sind Annahmen (OP-46)',None],
 ['STE-24','Steuerung und Physik','Regen','Meldet das Wettermodell Niederschlag, regnet es unter und in der untersten Wolkenschicht; über ihrer Obergrenze nicht. Im Regen sammelt sich Wasser auf der Hülle, bis ein Film von 0,1 mm gesättigt ist (rund 48 kg); je stärker der Regen, desto schneller. Das Wasser zählt zur Masse des Ballons. Nach dem Regen tropft und trocknet es in rund 30 Minuten ab. Regen kühlt die Hülle auf Lufttemperatur. Die Anzeige nennt Regen und Wasserlast; beim Einsetzen kommt eine Meldung. Kondenswasser ist nicht nachgebildet.','Im Regen von 2 mm/h trägt die Hülle nach 5 Minuten rund 16 kg und nach 30 Minuten rund 44 kg Wasser; über der Wolke bleibt sie trocken; danach nimmt die Last wieder ab.','Muss','Umgesetzt',N+'; Werte sind Annahmen (OP-46); Kondenswasser zurückgestellt',None]]
i=idx(rows,'MET-08')+1
rows.insert(i,['MET-09','Wetter','Wolken, Niederschlag und Tageszeit','Mit echtem Wetter lädt die App zusätzlich die Bedeckung je Druckfläche, die Bedeckung in drei Stockwerken als Ersatz und den Niederschlag. Daraus folgen die Wolkenschichten im Vertikalprofil, die Wolkendecke für die Strahlung und der Regen. Lehnt der Dienst die erweiterte Abfrage ab, lädt die App das Wetter ohne Wolken und Regen. Die Tageszeit ist die echte Uhrzeit am Startplatz. Die Simulation spielt immer an einem Sommermorgen ohne Wolken und Regen (21. Juni, 7 Uhr Sonnenzeit).','Das Startfenster nennt Bewölkung und Niederschlag; das Vertikalprofil zeigt die Wolkenschicht; in der Simulation steht die Sonne beim Start rund 27° hoch.','Muss','Umgesetzt',N+'; am echten Dienst noch nicht geprüft',None])
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.2',BLD)

ws=wb['Parameter']; ds=[copy.copy(ws.cell(40,c)._style) for c in range(1,5)]; r0=ws.max_row+1
P=[('Hülle: Fläche (Kugel mit 1000 m³)',484,'m²','Gerechnet; Durchmesser 12,4 m'),('Hülle: Flächengewicht',279,'g/m²','Gerechnet aus 135 kg Hülle und 484 m² (STE-23)'),
 ('Hülle: aufgenommener Anteil des Sonnenlichts',0.25,'–','Annahme Claude für weißes PVC-Gewebe; weiße Anstriche 0,2 bis 0,3 (OP-46)'),('Hülle: Emissionsgrad im Infraroten',0.9,'–','Annahme Claude; Kunststoffe rund 0,9; vorher 0,85 (OP-46)'),
 ('Himmel: Emissionsgrad bei klarem Himmel am Boden',0.74,'–','Annahme Claude; nimmt mit der Wurzel des Luftdrucks nach oben ab; Wolken darüber heben ihn auf 1'),
 ('Wolkendecke: durchgelassenes Sonnenlicht',0.3,'–','Annahme Claude; als Streulicht'),('Rückstrahlung: Boden / Wolkendecke darunter','0,2 / 0,6','–','Annahme Claude'),
 ('Regen: größte Wasserlast auf der Hülle',48,'kg','Annahme Claude: Film von 0,1 mm auf 484 m² (STE-24, OP-46)'),('Regen: Abtrocknen',30,'min','Annahme Claude: Zeitkonstante (STE-24)'),
 ('Simulation: Datum und Uhrzeit','21. Juni, 7 Uhr','Sonnenzeit','Vorgabe Nutzer: immer morgens bei Tag, ohne Wolken und Regen; Datum Vorschlag Claude (MET-09)'),('GPS: Dauer der Mittelung',10,'s','Vorgabe Nutzer 09.10.2026 (SPI-01)'),
 ('Gerechnet: Gas gegenüber Luft, klar, 500 m, mittags / 7 Uhr / nachts','+9,6 / +4,3 / −6,3','K','Ergebnis des Modells; je K rund 4 kg Auftrieb bei schlaffer Hülle')]
for k,v in enumerate(P):
    for c,x in enumerate(v,1): ws.cell(r0+k,c).value=x; ws.cell(r0+k,c)._style=copy.copy(ds[c-1])
    ws.cell(r0+k,2).number_format='General'

ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-01')]; r[6]=r[6]+'; der große Countdown zählt 10 Sekunden herunter und nennt den Grund'
r[11]=(r[11] or '')+f' Build {BLD}: Countdown am Rechner mit nachgestelltem Empfänger geprüft.'
tid=max(int(r[0][2:]) for r in rows)+1
rows.append([f'T-{tid}','Software','STE-23, STE-24, MET-09, R-19','Sonne, Nacht und Regen','Handy mit Netz; ein Tag mit Wolken oder Regen; Simulation','In der Simulation starten und 30 Minuten in rund 300 m fahren; mit echtem Wetter tagsüber, in der Dämmerung und nachts fahren; bei Regen unter die Wolke und darüber fahren; Hilfe-Anzeige für die Gastemperatur einschalten','Die Anzeige nennt Sonnenhöhe, Nacht, Wolken oder Regen; tagsüber ist das Gas wärmer als die Luft und der Ballon trägt besser, über den Sonnenuntergang braucht er Ballast, nachts ist das Gas kälter; im Regen wächst die Wasserlast und der Ballon sinkt, über der Wolke nicht; das Verhalten wirkt aus der Praxis glaubhaft','Handy, Hilfe-Anzeige','Mehrere Tageszeiten','2 Steuerung und Physik','Am Rechner geprüft',f'Build {BLD}: Sonnenstand geprüft (21.6. mittags 61,8°, 21.12. mittags 15°). Klarer Himmel, 500 m: Gas mittags +9,6 K, um 7 Uhr +4,3 K, nachts −6,3 K; unter Wolken nachts +1,3 K. Regen 2 mm/h: 16 kg nach 5 min, 44 kg nach 30 min, über der Wolke 0 kg. Wolken und Regen aus nachgestelltem Wetter erscheinen in Startfenster, Profil und Anzeige; abgelehnte Abfrage fällt auf Wetter ohne Wolken zurück. Am Handy offen: echter Dienst, Glaubwürdigkeit der Werte.','09.10.2026','Claude (Rechner)'])
r=rows[idx(rows,'T-06')]; r[11]=(r[11] or '')+f' Build {BLD}: Mit der Morgensonne der Simulation hält sich der Ballon nach 31 kg Abwurf vier Stunden zwischen 200 und 330 m, weil das Gas sich weiter erwärmt.'
put(ws,5,rows,14,4)

ws=wb['Risiken']; rows=rows_of(ws,6,15)
rows.append(['R-19','Spielgefühl','Sonne, Nachtabkühlung und Regen wirken zu stark oder zu schwach','Stoffwerte der Hülle, Strahlung des Himmels und Wasserlast sind Annahmen ohne Messwerte','Ballastverbrauch weicht von der Praxis ab; Vergleich mit echten Fahrten (Meisterschaft) hinkt',3,3,None,None,'Werte an einer Stelle einstellbar; Abgleich mit der Erfahrung des Nutzers und mit echten Fahrten',f'Werte nachstellen',f'T-{tid}','Nutzer, Claude','Offen','STE-23, STE-24'])
def fixR(v,r): v[7]=f'=F{r}*G{r}'; v[8]=f'=IF(H{r}>=15,"hoch",IF(H{r}>=8,"mittel","niedrig"))'
put(ws,6,rows,15,5,fixR)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
rows+=[['OP-46','Physik','Passen die gerechneten Wirkungen zur Praxis? Klarer Himmel in 500 m: Gas mittags rund 10 K wärmer und nachts rund 6 K kälter als die Luft, das sind rund +39 und −25 kg Auftrieb; über den Sonnenuntergang kostet das Halten der Höhe rund 55 kg Sand. Regen: bis 48 kg Wasser auf der Hülle. Annahmen dahinter: Hülle nimmt 25 % des Sonnenlichts auf, Emissionsgrad 0,9, Himmel 0,74, Wasserfilm 0,1 mm, Abtrocknen in 30 Minuten.','Werte bestätigen / abschwächen / verstärken','Nutzer','2','A',None,None,'Offen'],
       ['OP-47','Physik','Kondenswasser und Tau, Nebel am Boden und Schnee sind nicht nachgebildet; Niederschlag zählt immer als Regen. Die Wolkenbedeckung je Höhe schätzt das Wettermodell aus der Luftfeuchte. Soll davon später etwas ergänzt werden?','So lassen / Kondenswasser ergänzen / Schnee unterscheiden','Nutzer','2','C','Kondenswasser vorerst vernachlässigt (Nutzer 09.10.2026)',None,'In Klärung']]
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.27','09.10.2026',f'Build {BLD}: Sonne und Nachtabkühlung im Wärmehaushalt (STE-23), Regen mit Wasserlast (STE-24), Wolken, Niederschlag und Tageszeit aus dem Wetter, Simulation als Sommermorgen (MET-09); Anzeige von Himmel und Wasserlast, Wolken im Vertikalprofil (STE-08, STE-17). Trimmballast auf die Gastemperatur am Start abgestimmt. GPS mittelt 10 Sekunden mit großem Countdown (SPI-01). Neue Parameter, Test T-{tid}, Risiko R-19, offene Punkte OP-46 und OP-47.','Claude'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert T-%d'%tid)
