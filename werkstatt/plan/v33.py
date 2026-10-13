import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.33.xlsx").replace("0.18.xlsx","0.32.xlsx").replace("Version 0.19","Version 0.33").replace("Nutzer 08.10.2026","Nutzer 09.10.2026"))
BLD='2026-10-09.9'
ws=wb['Deckblatt']; ws['B4']='0.33'; ws['B6']=ws['B6'].value.replace('Build 2026-10-09.8',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,9)
r=rows[idx(rows,'MEI-03')]; r[3]=r[3].replace('ein roter Streifen von 11 km Breite: die 10 km Abstand der Regel und 1 km Reserve. Wer den Streifen nicht berührt, ist sicher.','ein roter Rand von 10 km Breite; er ist genau der Abstand der Regel. Wer den Rand berührt, ist disqualifiziert.')+' Die offizielle Restricted Area des Rennens ist zurückgestellt, bis die Daten vorliegen (OP-57).'; r[7]=(r[7] or '')+'; '+N+': Rand 10 km, rot'
r=rows[idx(rows,'MEI-06')]; r[3]=(r[3] or '')+' Entscheidung 09.10.2026: Bis die offiziellen Spuren vorliegen, zeigt die App nur die offiziellen Landepunkte.'; r[7]=(r[7] or '')+'; '+N
r=rows[idx(rows,'MEI-11')]; r[7]=(r[7] or '')+'; '+N+': 17 weitere Suchläufe kamen mit dem Wetterpaket (Raster 1° × 1,5°) nicht über rund 1500 km an Land; Finnland, der Landeort des Siegers, war darin nicht erreichbar'
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,9,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-09.8',BLD)
ws=wb['Parameter']
for i in range(1,ws.max_row+1):
    if ws.cell(i,1).value and str(ws.cell(i,1).value).startswith('Meisterschaft: Reserve'): ws.cell(i,2).value=0; ws.cell(i,4).value='Entscheidung Nutzer 09.10.2026: der rote Rand ist genau 10 km breit (MEI-03)'
ws=wb['Tests']; rows=rows_of(ws,5,14)
r=rows[idx(rows,'T-37')]; r[6]=r[6].replace('der Streifen ist rund 11 km breit','der rote Rand ist 10 km breit'); r[11]=(r[11] or '')+f' Build {BLD}: Rand auf 10 km gesetzt.'
for k,t in (('T-08',f' Unterer Rand: vom Nutzer am 09.10.2026 als in Ordnung gemeldet (OP-44).'),('T-22',f' Nutzer 09.10.2026: Start „viel besser jetzt“ (OP-52).')):
    r=rows[idx(rows,k)]; r[11]=(r[11] or '')+t
put(ws,5,rows,14,4)
ws=wb['Offene_Punkte']; rows=rows_of(ws,5,10)
def close(k,txt,st='Geschlossen'):
    r=rows[idx(rows,k)]; r[7]=((r[7]+' ') if r[7] else '')+txt; r[8]='09.10.2026'; r[9]=st
close('OP-44','Passt; lag wohl am Handy des Nutzers.')
close('OP-52','Viel besser seit Build 2026-10-09.6.')
close('OP-34','Die App zieht voraussichtlich mit dem nächsten Build auf das Synology-NAS um; das Paket ist dafür vorbereitet (LIESMICH mit Schritten für Synology). Eine gemeinsame Bestenliste ist damit noch nicht entschieden.','In Klärung')
rows+=[['OP-57','Meisterschaft','Offizielle Restricted Area und offizielle Spuren der Teams: zurückgestellt, der Nutzer besorgt die Daten in den nächsten Monaten. Bis dahin: nur die offiziellen Landepunkte, gesperrte Länder mit rotem 10-km-Rand. Die echten Teams landeten rund 35 bis 70 km vor den Landesgrenzen; die offizielle Fläche war also größer als im Spiel.','Daten liegen vor / bleibt zurückgestellt','Nutzer','7','C',None,None,'Offen'],
       ['OP-58','Meisterschaft','Feineres Wetterpaket für die Vorschau: Mit dem Raster 1° × 1,5° ist die Fahrt des Siegers nach Finnland nicht nachzurechnen. Soll die App ein feineres Paket (0,25°) entlang der Route laden, damit die KI-Route neu gerechnet werden kann?','Feineres Paket laden / so lassen','Nutzer','7','C',None,None,'Offen']]
put(ws,5,rows,10,4)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.33','09.10.2026',f'Build {BLD}: roter Rand um die gesperrten Länder genau 10 km (MEI-03). Geschlossen: OP-44, OP-52. Zurückgestellt: offizielle Restricted Area und Spuren (OP-57). Neu: OP-58 feineres Wetterpaket. Paket mit Anleitung für das Synology-NAS (OP-34).','Nutzer 09.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
