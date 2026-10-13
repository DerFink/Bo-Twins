import openpyxl, copy
exec(open('v19.py').read().split("# ---------------- Deckblatt")[0].replace("0.19.xlsx","0.43.xlsx").replace("0.18.xlsx","0.42.xlsx").replace("Version 0.19","Version 0.43").replace("Nutzer 08.10.2026","Nutzer 11.10.2026"))
BLD='2026-10-11.2'
ws=wb['Deckblatt']; ws['B4']='0.43'; ws['B6']=ws['B6'].value.replace('Build 2026-10-11.1',f'Build {BLD}')
ws=wb['Anforderungen']; rows=rows_of(ws,5,10)
def st(key,status=None,txt=None,note=None):
    r=rows[idx(rows,key)]
    if status: r[6]=status
    if txt: r[3]=(r[3] or '')+' '+txt
    r[7]=(r[7] or '')+'; '+N+(': '+note if note else '')
st('REC-04','Umgesetzt','Seit Build '+BLD+' geht der Standort für den Wetterabruf nur auf zwei Nachkommastellen gerundet (rund 1 km) an den Wetterdienst. Für die Geländehöhen werden weiterhin die Punkte eines Rasters um den Ballon abgefragt.','„bitte so übernehmen“')
st('NFA-01','Umgesetzt','In der Meisterschaft werden nur noch die Randstreifen der Sperrzonen im Umkreis von 450 km gezeichnet.','Nutzer: Karte läuft flüssig; einzelne Rückmeldungen „ruckelt manchmal“, vermutlich in der Meisterschaft')
st('NFA-02','Erledigt',None,'Nutzer: Knöpfe sind am Handy gut zu treffen')
st('NFA-03','Umgesetzt','Im Zielfahren erscheint das Startfenster sofort; der Ort des Ziels wird im Hintergrund gesucht und nachgetragen. Fährt der Spieler vorher los, bleibt das Ziel unverändert. Die Rechnung selbst dauert am Rechner 5 bis 50 ms; gewartet wurde auf die Ortssuche im Netz.','Nutzer: Zielberechnung dauert etwas lange')
st('MEI-12',None,'Der Knopf ist seit Build '+BLD+' ausgeblendet; die Funktion bleibt im Code.','Vorgabe Nutzer')
st('MEI-11',None,None,'11.10.2026: fünf weitere Suchläufe mit Stundenschritten und 13 Höhenbändern, auch mit dem Landeort des Siegers als Ziel: beste Landung an Land weiterhin rund 1490 km in Estland; Finnland bleibt mit diesem Wetterpaket unerreichbar')
st('STE-29',None,'Unter 8 % Füllgrad hängt die Hülle als schmale, flatternde Wurst am Korb, im Kartensymbol und in der Seitenansicht.','Wunsch Nutzer')
st('INS-06',None,'Der Knopf steht unten in den Einstellungen neben „Neueste Version laden“.','Wunsch Nutzer')
def fixA(v,r): v[8]=f'=COUNTIF(Tests!$C$5:$C$300,"*"&A{r}&"*")'
put(ws,5,rows,10,4,fixA); ws['A2']=ws['A2'].value.replace('2026-10-11.1',BLD)
ws=wb['Tests']; rows=rows_of(ws,5,15)
r=rows[idx(rows,'T-21')]; r[10]='Bestanden'; r[11]=(r[11] or '')+' Nutzer 11.10.2026: Die Karte läuft flüssig, die Knöpfe sind am Handy gut zu treffen. Einzelne Rückmeldungen „ruckelt manchmal“ ohne Einzelheiten; Build '+BLD+' zeichnet in der Meisterschaft weniger.'; r[12]='11.10.2026'; r[13]='Nutzer'
put(ws,5,rows,15,4)
ab=wb['Abnahme_Basis']
for r in range(ab.max_row,4,-1):
    if ab.cell(r,1).value in ('T-21','REC-04','NFA-01','NFA-02','NFA-03'): ab.delete_rows(r)
ws=wb['Änderungen']; r=ws.max_row+1
for c,v in enumerate(['0.43','11.10.2026',f'Build {BLD}: Standort für das Wetter gerundet (REC-04 umgesetzt). Zielfahren startet ohne Wartezeit (NFA-03). Meisterschaft zeichnet nur nahe Sperrzonen (NFA-01). Fast leerer Ballon als flatternde Hülle (STE-29). „Was ist neu“ unten in den Einstellungen (INS-06); Wetterpaket-Knopf ausgeblendet (MEI-12). T-21 bestanden, NFA-02 erledigt. Route des Siegers erneut gesucht, ohne Erfolg (MEI-11).','Nutzer 11.10.2026'],1):
    ws.cell(r,c).value=v; ws.cell(r,c)._style=copy.copy(ws.cell(r-1,c)._style)
ws.row_dimensions[r].height=None; ws.auto_filter.ref=f'A3:D{r}'
wb.save(OUT); print('gespeichert',OUT)
